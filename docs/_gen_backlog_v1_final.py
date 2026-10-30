# -*- coding: utf-8 -*-
"""Génère docs/Backlog_V1_Final.xlsx — backlog complet V1 (Excel client × Figma × CDC)."""
from copy import copy
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

DOCS = Path(__file__).resolve().parent
SRC = DOCS / "Backlog_Lotissement_PerduesDeVue.xlsx"
OUT = DOCS / "Backlog_V1_Final.xlsx"

NAVY, TEAL, GREEN, ORANGE, RED, GOLD = "1B365D", "0F6C8C", "1F7A4D", "C05621", "9B2C2C", "8A6D1F"
GREY, WHITE, INK = "F7F5F0", "FFFFFF", "1A1A1A"
thin = Border(
    left=Side(style="thin", color="D6D1C4"),
    right=Side(style="thin", color="D6D1C4"),
    top=Side(style="thin", color="D6D1C4"),
    bottom=Side(style="thin", color="D6D1C4"),
)
fill_navy = PatternFill("solid", fgColor=NAVY)
fill_teal = PatternFill("solid", fgColor=TEAL)
fill_grey = PatternFill("solid", fgColor=GREY)
font_h = Font(name="Calibri", bold=True, color=WHITE, size=11)
font_b = Font(name="Calibri", size=10, color=INK)
wrap = Alignment(wrap_text=True, vertical="top")

APP = {
    "PDV-000": "Dans l’app (socle ; recette/prod à brancher)",
    "PDV-001": "Dans l’app",
    "PDV-002": "Dans l’app (audit_logs)",
    "PDV-100": "Dans l’app",
    "PDV-101": "Dans l’app (Filament /admin)",
    "PDV-102": "Dans l’app",
    "PDV-103": "Dans l’app (invitation + MDP, pas magic link)",
    "PDV-104": "Dans l’app (/r/{token})",
    "PDV-105": "Dans l’app (Filament)",
    "PDV-200": "Dans l’app",
    "PDV-201": "Dans l’app",
    "PDV-202": "Dans l’app (Confirmer / Mettre à jour / Annuler)",
    "PDV-203": "Dans l’app",
    "PDV-204": "Dans l’app",
    "PDV-205": "Dans l’app (OTP email/SMS côté opticien)",
    "PDV-300": "Dans l’app (format FR + 4 chiffres + 2 lettres)",
    "PDV-301": "Dans l’app (already_engraved)",
    "PDV-302": "Dans l’app",
    "PDV-303": "Dans l’app",
    "PDV-304": "Dans l’app (irréversible)",
    "PDV-305": "Dans l’app",
    "PDV-306": "Dans l’app (archive ; code jamais recyclé)",
    "PDV-307": "Dans l’app (marques + Autre marque + couleurs)",
    "PDV-400": "Dans l’app (saisie / scanette ; pas caméra)",
    "PDV-401": "Dans l’app (PII masquée si pas client du magasin)",
    "PDV-402": "Dans l’app (mail/SMS + rate-limit)",
    "PDV-403": "Dans l’app (dossiers en cours seulement)",
    "PDV-404": "Dans l’app (KPI À graver ; file max 5 ; Remettre disabled si 0)",
    "PDV-405": "Dans l’app (inbox magasin)",
    "PDV-406": "Dans l’app (Next paramètres compte ; users magasin = Filament)",
    "PDV-500": "Dans l’app (6 statuts Figma)",
    "PDV-501": "Dans l’app",
    "PDV-502": "Dans l’app",
    "PDV-503": "Dans l’app (adresse + lien Maps ; carte intégrée / relais = bloqué clés)",
    "PDV-504": "Dans l’app (infos, connexion Modifier, notifs, suppression compte)",
    "PDV-600": "Dans l’app (gabarit + logs)",
    "PDV-601": "Dans l’app",
    "PDV-602": "Dans l’app (sandbox ; provider SMS prod à confirmer)",
    "PDV-700": "En app — domicile + relais liste ; carte / pins Colissimo+MR bloqués",
    "PDV-700-pay": "Dans l’app (Stripe Checkout redirect, Figma 25:1894)",
    "PDV-701": "En app — bordereau texte ; PDF transporteur bloqué sans API",
    "PDV-702": "Dans l’app (tracking_ref interne tant que pas d’API transporteur ; pas d’ETA)",
    "PDV-800": "Dans l’app (Filament KPIs)",
    "PDV-801": "Dans l’app (Filament opticiens + journal notifs)",
    "PDV-802": "Dans l’app (Filament clients / paires archivées / restitutions)",
    "PDV-803": "Dans l’app (Filament relation users magasin — pas Next)",
    "PDV-900": "Dans l’app (Filament lots + CSV + association magasin)",
    "PDV-901": "En app — lots Filament ; parc magasin-only = arbitrage (génération à l’enregistrement encore possible)",
    "PDV-1000": "En app — socle Filament ; barème prod + Qonto HTTP = arbitrage / clés",
    "PDV-1001": "Dans l’app (Filament tarifs réseau) — montants de prod = arbitrage",
    "PDV-1002": "Dans l’app (Filament fiche opticien factu)",
    "PDV-1003": "Dans l’app (contrôle + préparation mois ; pas de blocage magasin)",
    "PDV-1004": "Bloqué — page + logs ; skipped sans clé, queued sans HTTP",
    "OUT-EC-04": "Out V1",
    "OUT-NO-03": "Out V1",
    "OUT-CAMERA": "Hors V1 (saisie + HID)",
    "OUT-MFA": "Hors V1",
    "AU-01": "Dans l’app",
    "AU-02": "Dans l’app (Filament /admin)",
    "AU-03": "Dans l’app",
    "AU-04": "Dans l’app (invitation + MDP, pas magic link)",
    "AU-05": "Dans l’app (/r/{token})",
    "OP-01": "Dans l’app (KPI À graver = travail gravure, pas stock lots)",
    "OP-02": "Dans l’app (saisie / scanette ; pas caméra)",
    "OP-03": "Dans l’app (PII masquée si pas client du magasin)",
    "OP-04": "Dans l’app (pas de menu Utilisateurs Next)",
    "OP-05": "Dans l’app (Filament)",
    "CL-01": "Dans l’app",
    "CL-02": "Dans l’app",
    "CL-03": "Dans l’app",
    "CL-04": "Dans l’app (OTP email/SMS)",
    "CL-05": "Dans l’app",
    "LU-01": "Dans l’app (format FR + 4 chiffres + 2 lettres)",
    "LU-02": "Dans l’app",
    "LU-03": "Dans l’app (Filament lots) — parc magasin-only = arbitrage",
    "LU-04": "Dans l’app (Filament fiche codes du lot)",
    "LU-05": "Dans l’app (KPI À graver)",
    "LU-06": "Dans l’app",
    "LU-07": "Dans l’app (archive/inactive ; code JAMAIS recyclé)",
    "RE-01": "Dans l’app",
    "RE-02": "Dans l’app",
    "RE-03": "Dans l’app",
    "RE-04": "Dans l’app",
    "RE-05": "Dans l’app (dossiers en cours seulement)",
    "RE-06": "Dans l’app",
    "RE-07": "Dans l’app (inbox magasin)",
    "EC-01": "Dans l’app (6 statuts Figma)",
    "EC-02": "Dans l’app (adresse + lien Maps ; carte intégrée partielle)",
    "EC-03": "Dans l’app",
    "EC-04": "Out V1",
    "NO-01": "Dans l’app",
    "NO-02": "Dans l’app (fusionné dans NO-01)",
    "NO-03": "Out V1",
    "NO-04": "Dans l’app (mail client)",
    "NO-05": "Dans l’app (gabarit + logs)",
    "BO-01": "Dans l’app (Filament KPIs)",
    "BO-02": "Partiel Filament (consultation activité, pas classement 12 mois)",
    "BO-03": "Dans l’app (Filament + blocage)",
    "BO-04": "Dans l’app (Filament clients)",
    "BO-05": "Dans l’app (Filament restitutions en cours)",
    "BO-06": "Dans l’app (rôles Sanctum + Filament users magasin)",
    "GR-01": "Dans l’app (export code À graver)",
    "SE-01": "Dans l’app",
    "SE-03": "Dans l’app (audit_logs)",
    "AR-01": "Dans l’app (env local ; recette/prod à brancher)",
    "ADD-01": "Dans l’app",
    "ADD-02": "Dans l’app",
    "ADD-03": "Dans l’app",
    "ADD-04": "Dans l’app",
    "ADD-05": "Dans l’app (suppression compte + infos + connexion)",
    "ADD-06": "En app — carte relais / pins bloqués (clés Google + Colissimo/MR)",
    "ADD-07": "En app — bordereau texte ; PDF transporteur bloqué",
    "ADD-08": "Dans l’app (Filament, pas Next)",
}

LOT_APP = {
    0: "Dans l’app",
    1: "Dans l’app",
    2: "Dans l’app (users magasin = Filament, pas Next)",
    3: "Dans l’app",
    4: "Dans l’app",
    5: "Dans l’app",
    6: "Dans l’app (SMS sandbox / provider prod à confirmer)",
    7: "En app — Stripe Checkout OK ; carte relais + PDF transporteur bloqués",
    8: "Dans l’app (Filament)",
    9: "Dans l’app (Filament lots) — génération magasin encore possible (arbitrage)",
    10: "En app — contrôle + préparation ; sync HTTP Qonto bloquée",
}

CDC_REQ = [
    ("CODE-01", "4.1", "Saisie manuelle code lunettes", "3–4", "Dans l’app", "Must"),
    ("CODE-02", "4.1", "Saisie manuelle code de retrait", "4", "Dans l’app", "Must"),
    ("CODE-03", "4.1", "Contrôle format", "3", "Dans l’app (FR + 4 chiffres + 2 lettres)", "Must"),
    ("CODE-04", "4.1", "Contrôle existence", "3–4", "Dans l’app", "Must"),
    ("CODE-05", "4.1", "Contrôle statut", "3–4", "Dans l’app", "Must"),
    ("CODE-06", "4.1", "Copie presse-papiers", "3", "Dans l’app (export / affichage)", "Should"),
    ("CODE-07", "4.1", "Scan caméra smartphone", "—", "Hors V1 (saisie + scanette HID)", "Plus tard"),
    ("CODE-08", "4.1", "Lecture auto QR / codes-barres", "—", "Hors V1", "Plus tard"),
    ("CODE-09", "4.1", "Douchettes physiques", "4", "Partiel (HID clavier)", "Should"),
    ("CODE-10", "4.1", "Équipements tiers additionnels", "—", "Hors V1", "Plus tard"),
    ("CODE-11", "4.1", "Reconnaissance auto au-delà du scan", "—", "Hors V1", "Plus tard"),
    ("CODE-12", "4.2", "Code alphanumérique", "3", "Dans l’app", "Must"),
    ("CODE-13", "4.2", "Longueur 10 caractères", "3", "Dans l’app (FR####AA = 8 ; format PO figé)", "Must — écart CDC 10 car. tranché Figma/PO"),
    ("CODE-14", "4.2", "Unicité globale", "3", "Dans l’app", "Must"),
    ("CODE-15", "4.3", "Création lot admin", "9", "Dans l’app (Filament)", "Must"),
    ("CODE-16", "4.3", "Volume du lot", "9", "Dans l’app", "Must"),
    ("CODE-17", "4.3", "Génération des codes", "9", "Dans l’app", "Must"),
    ("CODE-18", "4.3", "Export CSV graveur", "9", "Dans l’app", "Must"),
    ("CODE-19", "4.3", "Historique des lots", "9", "Dans l’app (Filament)", "Must"),
    ("CODE-20", "4.3", "Statut des codes d’un lot", "9", "Dans l’app (fiche lot)", "Must"),
    ("CODE-21", "4.3", "API graveur directe", "—", "Hors V1 (CSV d’abord)", "Plus tard"),
    ("CODE-22", "4.3", "Import auto vers graveur", "—", "Hors V1", "Plus tard"),
    ("GEO-01", "5", "Nom opticien dépositaire", "5", "Dans l’app", "Must"),
    ("GEO-02", "5", "Adresse dépositaire", "5", "Dans l’app", "Must"),
    ("GEO-03", "5", "Téléphone / horaires", "5", "Partiel (tél magasin si renseigné ; horaires non Figma)", "Should"),
    ("GEO-04", "5", "Lien Maps externe", "5", "Dans l’app", "Must"),
    ("GEO-05", "5", "Carte intégrée", "5+7", "Bloqué clés Google (relais) ; lien Maps magasin OK", "Bloqué"),
    ("GEO-06", "5", "Point de retrait le plus proche", "7", "Liste relais stub ; carte pins bloquée", "Bloqué"),
    ("GEO-07", "5", "Itinéraire interne", "5", "Hors V1 (mémo : Maps/Waze)", "Plus tard"),
    ("GEO-08", "5", "Géoloc temps réel", "5", "Prompt 1re connexion (non bloquant)", "Should"),
    ("TARIF-01", "11.2", "Prix dégressifs volume", "10", "Hors V1 — modèle unitaire / forfait / gravure posé", "Plus tard"),
    ("TARIF-02", "11.2", "Paliers de quantité", "10", "Hors V1", "Plus tard"),
    ("TARIF-03", "11.2", "Forfaits", "10", "Socle Filament (type forfait) ; barème prod = arbitrage", "Partiel"),
    ("TARIF-04", "11.2", "Abonnements", "10", "Hors V1", "Plus tard"),
    ("TARIF-05", "11.2", "Remises automatiques", "10", "Hors V1", "Plus tard"),
    ("TARIF-06", "11.2", "Promotions", "10", "Hors V1", "Plus tard"),
    ("TARIF-07", "11.2", "Tarif par type de code", "10", "Hors V1", "Plus tard"),
    ("TARIF-08", "11.2", "Tarif par canal de génération", "10", "Hors V1", "Plus tard"),
    ("TARIF-09", "11.2", "Règles commerciales complexes", "10", "Hors V1", "Plus tard"),
    ("FACT-01", "12.5", "Préparation données facturation", "10", "Dans l’app (volumes + montant indicatif)", "Must"),
    ("FACT-02", "12.5", "Affichage factures prêtes", "10", "Contrôle Filament (pas d’émission Qonto)", "Partiel"),
    ("FACT-03", "12.5", "Affichage factures en erreur", "10", "Dans l’app (SIRET, réseau, tarif, IBAN…)", "Must"),
    ("FACT-04", "12.5", "Réévaluation après correction", "10", "Dans l’app (relancer le contrôle)", "Must"),
    ("FACT-05", "12.5", "Transmission lot Qonto", "10", "Bloqué sans clés (queued, pas d’HTTP)", "Bloqué"),
    ("FACT-06", "12.5", "Statut technique transmission", "10", "Logs Filament", "Partiel"),
    ("FACT-07", "12.5", "Sync statuts Qonto", "10", "Bloqué", "Bloqué"),
    ("FACT-08", "12.5", "N° facture Qonto", "10", "Bloqué", "Bloqué"),
    ("FACT-09", "12.5", "PDF facture Qonto", "10", "Bloqué", "Bloqué"),
    ("FACT-10", "12.5", "Suivi paiements", "10", "Hors V1 / bloqué Qonto", "Plus tard"),
    ("FACT-11", "12.5", "Suivi impayés", "10", "Hors V1", "Plus tard"),
    ("FACT-12", "12.5", "Remontée CA", "10", "Hors V1", "Plus tard"),
    ("FACT-13", "12.5", "Tableaux de bord financiers", "10", "Hors V1", "Plus tard"),
    ("FACT-14", "12.5", "Rapprochement comptable", "10", "Hors V1", "Plus tard"),
    ("FACT-15", "12.5", "Exports comptables avancés", "10", "Hors V1", "Plus tard"),
]

CDC_CHAPITRES = [
    ("§1", "Objectifs métier", "Tous", "Cœur livré (association, perte, découverte, notif, code retrait, cycle). Facturation opticien = socle Lot 10."),
    ("§2.1", "Admin PDV", "8–10", "Filament : réseaux, opticiens, users, clients, lots, paires, restitutions, notifs, tarifs, contrôle factu, page Qonto."),
    ("§2.2", "Admin opticien (users magasin)", "8", "Filament relation Utilisateurs du magasin — pas Next (Figma sans menu)."),
    ("§2.3", "Utilisateur opticien", "2–4", "Métier Next sans gestion des comptes."),
    ("§2.4", "Client final", "5–7", "Liste, perte, inactive, restitution, notifs, envoi à distance."),
    ("§3", "Réseaux / opticien / héritage tarif", "8+10", "Fiche Filament ; héritage prix/fréquence posé ; barème prod = arbitrage."),
    ("§4", "Codes", "3+9", "Voir feuille 09_CDC_exigences CODE-*."),
    ("§5", "Cartographie", "5+7", "Voir GEO-*."),
    ("§6", "BO admin", "8–10", "Filament (aucun écran Figma)."),
    ("§7", "BO opticien + PII + envoi postal", "2–7", "Next Figma. PII dépositaire. Étiquette post-paiement."),
    ("§8", "Espace client + envoi §8.7", "5+7", "Figma Client. Stripe Checkout. Relais carte bloquée."),
    ("§9", "Cycle de vie", "3–7", "Vocabulaire Figma (pas CDC « restituée ») : a_graver, active, perdue, inactive, a_recuperer, livraison_demandee, en_livraison, archivee."),
    ("§10", "3 parcours codes", "3+9", "Lot graveur + saisie déjà gravé + génération à l’enregistrement (les 3 dans l’app)."),
    ("§11", "Tarification", "10", "Unitaire / forfait / gravure. TARIF avancés hors V1."),
    ("§12", "Facturation + Qonto + envoi client", "7+10", "Stripe client OK. Qonto HTTP bloqué."),
    ("§13", "Notifications", "6–8", "Email+SMS transactionnels + logs. Enrichis hors V1. NO-03 out."),
    ("§14", "Auth / sécu / MFA", "0–1", "Sanctum, isolation, jetons publics, immuabilité. MFA hors V1."),
    ("§15", "RGPD", "0+5", "Minimisation PII, CGU, audit. Politique conservation à figer avec PDV."),
]


def style_header(ws, cols):
    for col in range(1, cols + 1):
        cell = ws.cell(1, col)
        cell.fill = fill_navy
        cell.font = font_h
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = thin


def style_body(ws, max_row, cols):
    for r in range(2, max_row + 1):
        for c in range(1, cols + 1):
            cell = ws.cell(r, c)
            cell.font = font_b
            cell.alignment = wrap
            cell.border = thin
            if r % 2 == 0:
                cell.fill = fill_grey


def autosize(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 28
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def insert_app_column(ws):
    # Col F is Statut cadrage in original (6). Insert after it a "Statut app V1 final".
    # Headers row 1: find "Statut cadrage"
    header_idx = None
    for c in range(1, ws.max_column + 1):
        val = ws.cell(1, c).value
        if val and "Statut cadrage" in str(val):
            header_idx = c
            break
    if header_idx is None:
        header_idx = 6
    insert_at = header_idx + 1
    ws.insert_cols(insert_at)
    head = ws.cell(1, insert_at)
    head.value = "Statut app (V1 final)"
    head.fill = fill_teal
    head.font = font_h
    head.alignment = Alignment(wrap_text=True, vertical="center")
    head.border = thin
    id_col = 1
    for r in range(2, ws.max_row + 1):
        sid = str(ws.cell(r, id_col).value or "").strip()
        cell = ws.cell(r, insert_at)
        cell.value = APP.get(sid, "")
        cell.font = font_b
        cell.alignment = wrap
        cell.border = thin


def add_sheet(wb, title, headers, rows, widths, header_fill=fill_navy):
    if title in wb.sheetnames:
        del wb[title]
    ws = wb.create_sheet(title)
    for i, h in enumerate(headers, 1):
        cell = ws.cell(1, i, h)
        cell.fill = header_fill
        cell.font = font_h
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = thin
    for r, row in enumerate(rows, 2):
        for c, val in enumerate(row, 1):
            cell = ws.cell(r, c, val)
            cell.font = font_b
            cell.alignment = wrap
            cell.border = thin
            if r % 2 == 0:
                cell.fill = fill_grey
    autosize(ws, widths)
    last = get_column_letter(len(headers))
    tab = Table(displayName=title.replace(" ", "_")[:20] + "T", ref=f"A1:{last}{1 + len(rows)}")
    tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    try:
        ws.add_table(tab)
    except Exception:
        pass
    return ws


def main():
    wb = load_workbook(SRC)

    # Légende : append V1 final notes
    leg = wb["00_Legende"]
    start = leg.max_row + 2
    notes = [
        ("LIVRABLE V1 FINAL", "Ce fichier est le backlog complet : 48 stories Excel client + ajouts Figma (ADD) + CDC + PDV-*."),
        ("Arbitrage", "Mémos Figma > CDC > Excel client. Décisions : docs/00-cadrage-et-decisions.md"),
        ("Figma", "https://www.figma.com/design/7bIosD4oOxnQl8iKfHX9O6/Perdues-de-vue----Wireframes?node-id=1-2"),
        ("Client", "Page 1:535 (mobile 375)"),
        ("Opticien", "Page 7:1919 (desktop 1440)"),
        ("Admin", "Aucun wireframe — Filament Laravel"),
        ("Statut app", "Dans l’app · Partiel · En app (bloqué clés) · Out V1 · Plus tard"),
        ("Users magasin", "CDC §7.2 : Filament uniquement. Pas d’écran Next (absent Figma)."),
        ("Paiement envoi", "Stripe Checkout hébergé (redirect), Figma 25:1894 — pas de carte dans l’app."),
        ("Onglet ajouté", "09_CDC_exigences = CODE / GEO / TARIF / FACT. 10_CDC_chapitres = §1–15. 11_Synthese_V1."),
    ]
    for i, (k, v) in enumerate(notes):
        a = leg.cell(start + i, 1, k)
        b = leg.cell(start + i, 2, v)
        a.font = Font(name="Calibri", bold=True, color=WHITE, size=10)
        a.fill = fill_teal
        b.font = font_b
        b.alignment = wrap

    # Lotissement : add statut app
    lot = wb["01_Lotissement"]
    col = lot.max_column + 1
    h = lot.cell(1, col, "Statut app (V1 final)")
    h.fill = fill_teal
    h.font = font_h
    h.border = thin
    for r in range(2, lot.max_row + 1):
        lot_id = lot.cell(r, 1).value
        cell = lot.cell(r, col, LOT_APP.get(int(lot_id) if lot_id is not None else -1, ""))
        cell.font = font_b
        cell.alignment = wrap
        cell.border = thin

    insert_app_column(wb["02_Backlog"])

    add_sheet(
        wb,
        "09_CDC_exigences",
        ["ID CDC", "§", "Exigence", "Lot", "Statut app (V1 final)", "MoSCoW / note"],
        [list(x) for x in CDC_REQ],
        [14, 8, 48, 10, 55, 42],
    )
    add_sheet(
        wb,
        "10_CDC_chapitres",
        ["Chapitre CDC", "Sujet", "Lots", "Couverture V1 finale"],
        [list(x) for x in CDC_CHAPITRES],
        [14, 36, 12, 80],
    )

    synthese = [
        ("Stories Excel client", "48 (AU…AR) — toutes listées onglet 02_Backlog"),
        ("Ajouts Figma/CDC", "ADD-01…08 + PDV-700…702, 801, 900, 1001…1004 (08_Demandes_client)"),
        ("Écrans Figma inventoriés", "85 (onglet 03_Ecrans_Figma) — 0 écran Admin"),
        ("Mémos UX", "34 (onglet 05_Memos_UX) — source de vérité parcours"),
        ("Exigences CDC numérotées", f"{len(CDC_REQ)} (CODE/GEO/TARIF/FACT)"),
        ("Chapitres CDC", "§1 à §15 tous couverts ou explicitement hors/bloqués"),
        ("Cœur magasin (lots 0–6)", "Dans l’app"),
        ("Envoi à distance (lot 7)", "Dans l’app sauf carte relais + PDF API transporteur"),
        ("Admin / lots / factu (8–10)", "Filament dans l’app ; Qonto HTTP bloqué ; barème prod = arbitrage"),
        ("Out V1", "EC-04 inscription publique · NO-03 notif perte opticien d’origine · recycle code · MFA · caméra · TARIF avancés · FACT-10…15"),
        ("Bloquants client", "Clés Google Maps · contrats/API Colissimo + Mondial Relay · compte/clés Qonto · provider SMS prod · liste réseaux + unité tarif"),
    ]
    add_sheet(
        wb,
        "11_Synthese_V1",
        ["Item", "V1 finale"],
        synthese,
        [36, 90],
        header_fill=fill_teal,
    )

    # Move new sheets after 08
    order = [
        "00_Legende",
        "01_Lotissement",
        "02_Backlog",
        "03_Ecrans_Figma",
        "04_Menus_actions",
        "05_Memos_UX",
        "06_Hors_scope",
        "07_Chemin_demo",
        "08_Demandes_client",
        "09_CDC_exigences",
        "10_CDC_chapitres",
        "11_Synthese_V1",
    ]
    for i, name in enumerate(order):
        wb.move_sheet(name, offset=i - wb.sheetnames.index(name))

    wb.save(OUT)
    print("Wrote", OUT, "sheets", wb.sheetnames)


if __name__ == "__main__":
    main()
