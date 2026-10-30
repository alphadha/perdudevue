# -*- coding: utf-8 -*-
"""BacklogFinal.xlsx — une seule feuille, format atelier (ID, Epic, Type, US, AC, edge cases)."""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

DOCS = Path(__file__).resolve().parent
OUT = DOCS / "BacklogFinal.xlsx"
FIGMA = "https://www.figma.com/design/7bIosD4oOxnQl8iKfHX9O6/Perdues-de-vue----Wireframes"

NAVY, TEAL, GOLD, GREY, WHITE, INK = "1B365D", "0F6C8C", "8A6D1F", "F7F5F0", "FFFFFF", "1A1A1A"
ORANGE = "C05621"
thin = Border(
    left=Side(style="thin", color="D6D1C4"),
    right=Side(style="thin", color="D6D1C4"),
    top=Side(style="thin", color="D6D1C4"),
    bottom=Side(style="thin", color="D6D1C4"),
)
fill_navy = PatternFill("solid", fgColor=NAVY)
fill_teal = PatternFill("solid", fgColor=TEAL)
fill_gold = PatternFill("solid", fgColor=GOLD)
fill_grey = PatternFill("solid", fgColor=GREY)
fill_orange = PatternFill("solid", fgColor=ORANGE)
font_h = Font(name="Calibri", bold=True, color=WHITE, size=11)
font_title = Font(name="Calibri", bold=True, size=16, color=NAVY)
font_b = Font(name="Calibri", size=10, color=INK)
font_banner = Font(name="Calibri", bold=True, color=WHITE, size=12)
wrap = Alignment(wrap_text=True, vertical="top")
wrap_c = Alignment(wrap_text=True, vertical="center")

COLS = 9
HEADERS = [
    "ID",
    "Epic",
    "Type",
    "Résumé",
    "Description",
    "Critères d'acceptances",
    "Edge cases",
    "Version",
    "Dépendance API",
]


def bullets(*items):
    return "\n".join(f"- {x}" for x in items)


def link(node):
    return f"{FIGMA}?node-id={node.replace(':', '-')}"


# Stories : (id, epic, type, resume, description, ac, edges, version, api)
# Bandeaux historiques ignorés — l’ordre final = FLOW (parcours d’usage).

ROWS = [
    (
        "FO-C-01", "Connexion", "US", "Connexion client — email + mot de passe",
        "Permettre au client déjà invité par son opticien d’accéder à son espace « Mes lunettes ».\nMaquette : Connexion " + link("1:602") + "\nCDC §8 / §14. Copy Figma : pas d’inscription autonome (« rapprochez-vous de votre opticien »).",
        bullets(
            "L’utilisateur saisit un email et un mot de passe.",
            "Le bouton de connexion n’est actif que si les champs requis sont remplis.",
            "Identifiants valides + compte activé → page d’accueil Mes lunettes (1:548).",
            "Identifiants incorrects → message générique (pas « email inconnu »).",
            "Compte non activé (invitation non terminée) → accès refusé, message générique.",
            "Lien « mot de passe oublié » ouvre le flux reset (1:982).",
            "Toggle affichage du mot de passe selon Figma.",
        ),
        bullets(
            "Clic connexion avec champs vides → validation champs + message.",
            "Email mal formé → « Adresse e-mail invalide ».",
            "Email / MDP faux → même message, aucune énumération d’emails.",
            "Tentative répétée → limitation des essais (sécurité).",
            "Utilisateur opticien qui utilise cet écran → pas d’accès espace client.",
        ),
        "V1", "Non",
    ),
    (
        "FO-C-02", "Connexion", "US", "Mot de passe oublié client (code à 6 chiffres)",
        "Permettre au client de réinitialiser son mot de passe par email : demande → code 6 chiffres → nouveau MDP.\nMaquettes : 1:982, 1:1001, 1:1038.",
        bullets(
            "Saisie d’un email → message de confirmation identique que l’email existe ou non.",
            "Code à 6 chiffres, à usage unique, durée de vie limitée.",
            "Nouveau mot de passe avec confirmation + règles de complexité (à figer, min. 8 car. proposé).",
            "Après succès → écran de connexion.",
        ),
        bullets(
            "Code expiré / déjà utilisé → refus, possibilité de renvoyer.",
            "Code erroné → message, sans indiquer si l’email est connu.",
            "Les deux MDP ne correspondent pas → blocage.",
        ),
        "V1", "Email transactionnel",
    ),
    (
        "FO-C-03", "Connexion", "US", "Activation compte client — invitation + mot de passe + CGU",
        "Le client créé par l’opticien reçoit un email d’accès. Infos déjà remplies ; il ne définit que son mot de passe et accepte CGU / politique de confidentialité.\nExcel AU-04 parlait de magic link sans compte : Figma = set password (décision à valider D-02).\nCDC §13 invitation / activation.",
        bullets(
            "Lien d’invitation valide → écran pose du MDP + cases CGU.",
            "Sans acceptation CGU : bouton de validation inactif ou refus.",
            "Succès → espace client, infos perso préremplies.",
            "Invitation à usage unique.",
        ),
        bullets(
            "Lien expiré / déjà consommé → message d’invalidité, pas d’accès.",
            "Lien falsifié → page introuvable.",
            "Client jamais créé par un opticien → pas de parcours d’inscription publique (voir FO-C-04).",
        ),
        "V1", "Email transactionnel",
    ),
    (
        "FO-C-04", "Connexion", "US", "Hors V1 proposé — inscription publique autonome (EC-04)",
        "Excel EC-04 : un particulier crée son compte sans opticien. Copy Figma inverse. CDC : le compte client est créé par le magasin.",
        bullets(
            "Si Validé hors V1 : aucun formulaire d’inscription grand public.",
            "L’écran connexion oriente vers l’opticien / support.",
        ),
        bullets(
            "Si le client exige EC-04 en V1 : nouvelle maquette à produire (absente Figma).",
        ),
        "Hors V1", "Non",
    ),
    (
        "FO-C-05", "Connexion", "US", "Page publique « lunettes retrouvées » (sans connexion)",
        "Le client ouvre le lien reçu par email et voit, sans session, les infos pour récupérer SA paire, et un accès pour se reconnecter.\nMaquette : 1:776. CDC §8.3 / §14 jeton public. Excel AU-05 / EC-02.",
        bullets(
            "Jeton non séquentiel, durée de vie limitée.",
            "Affiche une seule paire : magasin dépositaire, adresse, lien carto, code de restitution, statut.",
            "CTA pour ouvrir / se reconnecter au compte.",
            "Pas d’autres paires du client sur cette page.",
        ),
        bullets(
            "Jeton falsifié / expiré → 404 générique.",
            "Paire déjà remise → page cohérente (plus en attente) ou lien invalidé (à figer).",
        ),
        "V1", "Non",
    ),
    (
        "FO-C-10", "Espace client", "US", "Liste Mes lunettes — ordre et badges Figma",
        "Le client consulte toutes ses paires. Ordre : À récupérer → En livraison → Livraison demandée → Perdue → Active → Inactive.\nMaquette : Page d’accueil 1:548. Excel EC-01 (vocabulaire Excel Enregistrée/Restituée remplacé par Figma). CDC §8.1 / §9.",
        bullets(
            "Chaque carte affiche type + marque, statut, actions Figma.",
            "Le statut « À graver » n’apparaît pas côté client.",
            "Clic carte → fiche du statut correspondant.",
            "Liste vide → état vide Figma.",
        ),
        bullets(
            "Client sans paire.",
            "Client avec paires chez plusieurs opticiens : il voit toutes SES paires, pas celles d’un autre compte.",
        ),
        "V1", "Non",
    ),
    (
        "FO-C-11", "Espace client", "US", "Fiche paire Active",
        "Détail d’une paire utilisable : code assigné, type, marque, historique (plus récent en premier), actions Déclarer perdue / inactive.\nMaquette : 1:624.",
        bullets(
            "Titre = type + marque (ex. Lunettes Vue - Gucci).",
            "Code assigné visible, copiable.",
            "Actions conformes au menu Figma du statut Active.",
        ),
        bullets(
            "Paire d’un autre client → inaccessible.",
            "Paire déjà Perdue : cette fiche n’est plus Active.",
        ),
        "V1", "Non",
    ),
    (
        "FO-C-12", "Espace client", "US", "Déclarer une paire perdue (réversible)",
        "Le client signale une perte temporaire, avec confirmation.\nPop-in 1:906, suite 1:654, fiche Perdue 1:829. CDC §8.2. Excel RE-01.",
        bullets(
            "Confirmation obligatoire avant validation.",
            "Statut → Perdue, horodatage, historique.",
            "Possibilité d’annuler la déclaration → retour Active.",
            "Toaster / copy Figma.",
        ),
        bullets(
            "Fermer le pop-in sans confirmer → aucune modification (toaster abandon).",
            "Paire Inactive / archivée : action absente.",
            "Paire déjà À récupérer : pas de nouvelle déclaration de perte.",
        ),
        "V1", "Non",
    ),
    (
        "FO-C-13", "Espace client", "US", "Déclarer inactive / hors d’usage (définitif)",
        "Le client sort la paire du cycle (cassée, perdue définitive).\nPop-in 17:1979 / 4:937, fiches 4:1430 / 4:1584. CDC §8.6 / §2.4.",
        bullets(
            "Double confirmation, action définitive.",
            "Paire Inactive visible avec date / motif.",
            "Code assigné définitivement inutilisable (jamais recyclé).",
            "Plus de perte, plus de découverte, plus de notifs liées.",
        ),
        bullets(
            "Abandon pop-in → inchangé.",
            "Tentative de signaler retrouvée sur une inactive → refus métier.",
        ),
        "V1", "Non",
    ),
    (
        "FO-C-14", "Espace client", "US", "Supprimer définitivement une paire",
        "La paire disparaît des listes client et opticien ; conservée en BO (archivée) pour traçabilité.\nPop-in 1:862. CDC / mémo suppression.",
        bullets(
            "Confirmation irréversible.",
            "Plus visible ni chez le client ni chez l’opticien.",
            "Le code n’est pas remis à disposition.",
        ),
        bullets(
            "Paire en cours de restitution : à interdire ou clôturer d’abord (à figer).",
        ),
        "V1", "Non",
    ),
    (
        "FO-C-15", "Espace client", "US", "Fiche À récupérer + localisation magasin",
        "Après découverte : le client voit où récupérer, le code de restitution, le choix magasin vs envoi.\nMaquettes 1:693 / 1:738 / 17:3441, pop-in carte 1:884. CDC GEO-01 à 04, §8.3.",
        bullets(
            "Nom, adresse, téléphone si dispo, lien « Ouvrir dans Maps / Waze ».",
            "Code de restitution affiché.",
            "Option se faire envoyer (FO-C-20).",
            "Prompt géoloc à la 1re connexion (non bloquant) — GEO-08.",
        ),
        bullets(
            "Horaires magasin absents → ne pas afficher de bloc vide trompeur (GEO-03 à confirmer).",
            "Refus de géoloc → écran utilisable quand même.",
        ),
        "V1", "Service cartographique (lien externe)",
    ),
    (
        "FO-C-16", "Espace client", "US", "Fiches Livraison demandée et En livraison",
        "Le client suit l’envoi après paiement.\nMaquettes 46:2645, 46:2434, 32:3292. CDC §8.7 / §9.",
        bullets(
            "Statuts : en attente d’envoi, expédiée / en livraison.",
            "N° de suivi transporteur s’il est disponible.",
            "Pas de nouvelle demande d’envoi sur la même restitution ouverte.",
        ),
        bullets(
            "Paiement non confirmé → la paire reste À récupérer (pas Livraison demandée).",
            "ETA : ne pas afficher de date promise sans contrat transporteur (mémo 63:2111).",
        ),
        "V1", "Transporteur (suivi)",
    ),
    (
        "FO-C-20", "Envoi à distance", "US", "Choix du mode : magasin / domicile / point relais",
        "Alternative si le magasin est trop loin.\nÉcran 18:1603. CDC parcours de demande §8.7.",
        bullets(
            "Choix mutuellement exclusifs selon Figma.",
            "Retour possible vers la fiche À récupérer.",
            "Magasin = parcours retrait (pas de paiement envoi).",
        ),
        bullets(
            "Paire pas À récupérer → tunnel inaccessible.",
            "Abandon en cours → toaster « modifications non prises en compte ».",
        ),
        "V1", "Non",
    ),
    (
        "FO-C-21", "Envoi à distance", "US", "Choix d’un point relais (plus proche)",
        "Le client sélectionne un relais Colissimo ou Mondial Relay (CDC : transporteurs à confirmer).\nÉcran 21:1717. GEO-06.",
        bullets(
            "Liste / carte des points proches.",
            "Sélection obligatoire d’un relais avant récap.",
            "Pins / carte intégrée (GEO-05) : exigence CDC/Figma à valider avec le prestataire carto.",
        ),
        bullets(
            "Aucun relais à proximité → message + bascule domicile possible.",
            "Changement d’adresse / géoloc → rafraîchir la liste.",
        ),
        "V1", "Colissimo / Mondial Relay + cartographie",
    ),
    (
        "FO-C-22", "Envoi à distance", "US", "Récapitulatif, adresse, consentement, prix",
        "Le client confirme l’adresse d’expédition, le consentement (nom + adresse uniquement pour l’étiquette) et le montant.\nÉcran 18:1665. Figma 2,99 € — montant à figer (D-11).",
        bullets(
            "Case consentement obligatoire, texte informé CDC.",
            "Adresse complète obligatoire pour domicile.",
            "Affichage du montant TTC du service (expédition + traitement).",
            "Sans paiement validé : demande non confirmée.",
        ),
        bullets(
            "Consentement non coché → impossible de payer.",
            "Adresse incomplète → validation champ par champ.",
        ),
        "V1", "Non",
    ),
    (
        "FO-C-23", "Envoi à distance", "US", "Paiement du service d’envoi (encaissé par la plateforme)",
        "Paiement en ligne client ; l’opticien ne gère ni paiement ni affranchissement.\nÉcran tunnel Figma 25:1894, échec 25:1767, succès → 46:2645. CDC §12.6. Prestataire (Stripe sur maquette) à confirmer.",
        bullets(
            "Redirection vers le paiement hébergé (pas de saisie CB dans l’app métier).",
            "Succès → statut Livraison demandée + notif opticien + génération étiquette.",
            "Échec / abandon → paire reste À récupérer, message Figma.",
            "Justificatif / facture simplifiée client (CDC).",
            "Indépendant de la facture opticien (prix par code).",
        ),
        bullets(
            "Retour navigateur / timeout → pas de statut livré.",
            "Double clic payer → un seul encaissement.",
            "Montant 2,99 € Figma vs tarif réel transporteur : décision D-11.",
        ),
        "V1", "Paiement en ligne (maquette Stripe)",
    ),
    (
        "FO-C-24", "Espace client", "US", "Carte intégrée et itinéraire voiture dans l’app (GEO-05 / GEO-07)",
        "CDC + mémo géoloc : carte dans l’application, itinéraire interne = voiture uniquement ; autres modes via Maps/Waze.",
        bullets(
            "Carte centrée sur le client / le magasin.",
            "Itinéraire voiture interne si l’exigence est retenue en V1.",
            "Liens Maps / Waze toujours disponibles.",
        ),
        bullets(
            "Géoloc refusée → carte par défaut (adresse magasin).",
        ),
        "V1.1", "Service cartographique intégré",
    ),
    (
        "FO-C-30", "Paramètres client", "US", "Hub paramètres, profil, connexion, MDP, téléphone, notifications",
        "Le client gère son compte.\nMaquettes : 1:939, 1:1193, 1:1065, 1:1109–1:1150, 4:1741, 1:1242, pop-in déconnexion 1:919. CDC §8.5. Excel ADD-05 / SE-01.",
        bullets(
            "Infos perso consultables / modifiables.",
            "Changement email → OTP email avant persistance.",
            "Changement téléphone → OTP SMS avant persistance.",
            "Préférences email et SMS séparées.",
            "Modification MDP avec MDP actuel.",
            "Déconnexion via pop-in Figma.",
        ),
        bullets(
            "OTP erroné / expiré → pas de changement d’identité.",
            "Email déjà utilisé par un autre compte → refus.",
        ),
        "V1", "Email + SMS (OTP)",
    ),
    (
        "FO-C-31", "Paramètres client", "US", "Suppression du compte client",
        "Pop-in 1:929. CDC §15 politique de conservation à figer (D-19).",
        bullets(
            "Confirmation explicite.",
            "Conséquences affichées (accès, paires, restitutions en cours).",
            "Traçabilité BO selon politique RGPD.",
        ),
        bullets(
            "Restitution en cours : interdire ou clôturer (à figer).",
            "Abandon pop-in → compte inchangé.",
        ),
        "V1", "Non",
    ),
    (
        "FO-O-01", "Connexion magasin", "US", "Connexion opticien email + mot de passe",
        "Accès à l’espace magasin. Maquette 41:2323. Excel AU-01. CDC §14. Isolation par établissement.",
        bullets(
            "Email + MDP → tableau de bord 7:1920.",
            "Erreur générique si échec.",
            "Pas d’auto-inscription : copy contact support@perduesdevue.fr.",
            "Un user magasin n’accède qu’à SON opticien.",
        ),
        bullets(
            "Client qui se trompe d’espace → pas de données magasin.",
            "Admin PDV n’utilise pas cet écran (espace distinct).",
            "Compte magasin désactivé → refus.",
        ),
        "V1", "Non",
    ),
    (
        "FO-O-02", "Connexion magasin", "US", "Reset MDP + invitation opticien (vérif mail)",
        "Maquettes 41:2522, 41:2582, 41:2642. L’admin PDV crée le magasin et envoie les accès (Excel OP-05).",
        bullets(
            "Reset : code / lien puis nouveau MDP.",
            "Première connexion après invitation : pose du MDP.",
            "Messages sans énumération d’emails.",
        ),
        bullets(
            "Invitation expirée.",
            "Opticien inactif / bloqué.",
        ),
        "V1", "Email transactionnel",
    ),
    (
        "FO-O-10", "Espace opticien", "US", "Tableau de bord KPIs, actions rapides, file max 5",
        "Maquette 7:1920. Excel OP-01/02/04/LU-05. CDC §7.1.\nMenu Figma : Tableau de bord, Clients, Lunettes, Restitution, Notifications, Paramètres. Pas de menu Utilisateurs (users = BO, CDC §7.2).",
        bullets(
            "KPIs Figma dont file d’attente et « À graver ».",
            "Actions : créer client, enregistrer paire, signaler retrouvée, remettre.",
            "Bloc lunettes retrouvées : 5 paires max.",
            "CTA « Remettre des lunettes » actif seulement s’il existe ≥ 1 paire À récupérer, sinon disabled.",
            "Libellé « Client notifié » sur la file.",
        ),
        bullets(
            "File vide.",
            "Plus de 5 dossiers : le bloc n’en montre que 5, le reste sur Restitutions.",
            "Nom du client dans la file : uniquement s’il est déjà client de CE magasin (PII).",
        ),
        "V1", "Non",
    ),
    (
        "FO-O-11", "Espace opticien", "US", "Paramètres compte magasin (pas la gestion d’équipe)",
        "Maquettes 7:4995, 7:5162.",
        bullets(
            "Informations de connexion du user connecté.",
            "Pas de CRUD des collègues sur ces écrans Figma.",
        ),
        bullets(
            "User simple vs admin magasin : mêmes écrans compte, droits users ailleurs (BO-O-01).",
        ),
        "V1", "Non",
    ),
    (
        "FO-O-12", "Notifications magasin", "US", "Inbox / cloche opticien",
        "Maquette 14:1657. Excel RE-07. CDC §13 : envoi à distance payé, étiquette, etc.",
        bullets(
            "Liste des notifications du magasin.",
            "Clic « nouvel envoi demandé » → dossier Livraison demandée.",
            "Notification d’envoi seulement après paiement client confirmé.",
        ),
        bullets(
            "Autre magasin : 0 notif de ce dossier.",
        ),
        "V1", "Non",
    ),
    (
        "FO-O-20", "Gestion clients", "US", "Créer un client + email de bienvenue",
        "Pop-in 7:3230 / 41:3172, fiche 41:2907 / 7:2633. Excel CL-01/CL-05/NO-04. CDC §7.3.\nPrénom, nom, email, tél optionnel, consentements email/SMS séparés.",
        bullets(
            "Toaster succès puis redirection fiche client.",
            "Email d’invitation / bienvenue avec lien d’accès.",
            "Email unique global.",
        ),
        bullets(
            "Email invalide / champs requis vides → validation.",
            "Téléphone vide autorisé.",
            "Doublon → FO-O-21, pas un 2e compte.",
        ),
        "V1", "Email transactionnel",
    ),
    (
        "FO-O-21", "Gestion clients", "US", "Doublon : Confirmer / Mettre à jour / Annuler",
        "Pop-in 25:2301 + modifier fiche 77:2181. Mémo Figma critique.",
        bullets(
            "Si nom/prénom/email et/ou tél déjà associés : pop-in compte existant (infos masquées/affichées selon règle).",
            "Confirmer → rattacher le client à CET opticien, fin du flow.",
            "Mettre à jour → pop-in modifier fiche (OTP si email/tél) puis reprise de la création.",
            "Annuler → retour formulaire, champs conservés, encore modifiables.",
        ),
        bullets(
            "Client déjà rattaché à ce magasin → message adapté, pas de doublon de rattachement.",
            "Client uniquement chez un autre magasin : Confirmer le rattache ici sans exposer l’autre magasin inutilement.",
        ),
        "V1", "Email / SMS si OTP",
    ),
    (
        "FO-O-22", "Gestion clients", "US", "Liste, recherche, fiche, modifier, détacher",
        "Liste 7:2040, fiche 7:2633, pop-ins modifier / supprimer 7:3144, 7:3284. Excel CL-02/03/04. CDC §7.3.",
        bullets(
            "Recherche nom, email, autre critère pertinent.",
            "Fiche : paires + restitutions du client dans ce magasin.",
            "Changement email → OTP email ; tél → OTP SMS ; pas de persist avant OTP.",
            "Détacher : l’autre magasin conserve son rattachement.",
            "Enregistrer une paire depuis la fiche : client prérempli.",
        ),
        bullets(
            "Liste A ne contient pas un client uniquement chez B.",
            "Fiche d’un client non rattaché → inaccessible.",
            "OTP KO → identité inchangée.",
        ),
        "V1", "Email + SMS (OTP)",
    ),
    (
        "FO-O-30", "Lunettes & codes", "US", "Enregistrer une paire — génération automatique du code (À graver)",
        "Pop-in sans code 7:3354 (variantes 41:3251, 41:3808). CDC §10.3 / §7.4. Excel LU-01.\nTypes Figma : Vue / Soleil / Solaire de Vue. Marque autocomplete + Autre marque. Couleurs multi-select.\nFormat PO proposé : FR + 4 chiffres + 2 lettres (ex. FR5487OK) — CDC CODE-13 = 10 car. (décision D-01).\nL’association est facturable (CDC §12.2), pas la génération d’un lot graveur.",
        bullets(
            "Client obligatoire (autocomplete dashboard ; prérempli depuis fiche client).",
            "Un code unique est attribué, visible dans la fiche.",
            "Statut À graver.",
            "Association définitive, code jamais réutilisé.",
        ),
        bullets(
            "Client non sélectionné dans la liste → impossible d’enregistrer.",
            "Collision de code → nouvel essai automatique, unicité globale CODE-14.",
        ),
        "V1", "Non",
    ),
    (
        "FO-O-31", "Lunettes & codes", "US", "Enregistrer une paire — code déjà gravé",
        "Pop-in avec code 9:5459 (variantes). CDC §7.5 / §10.2.",
        bullets(
            "Contrôle format, existence/validité, non déjà associé.",
            "Statut Active dès l’association.",
            "Trace user, magasin, date. Événement facturable.",
        ),
        bullets(
            "Code inconnu / mauvais format → refus CODE-03/04.",
            "Code déjà associé → refus, pas d’overwrite.",
            "Code d’une paire inactive / archivée → refus.",
        ),
        "V1", "Non",
    ),
    (
        "FO-O-32", "Lunettes & codes", "US", "Indiquer comme gravée + export codes à graver",
        "Pop-in 46:2201, multi-select 41:4989. Excel GR-01.",
        bullets(
            "À graver → Active, irréversible, confirmation.",
            "Export CSV/TXT des codes sélectionnés (uniquement À graver).",
            "Checkboxes des autres statuts désactivées sur ce multi-select.",
        ),
        bullets(
            "Abandon pop-in → statut inchangé.",
            "Export vide si aucune sélection.",
        ),
        "V1", "Non",
    ),
    (
        "FO-O-33", "Lunettes & codes", "US", "Liste et 7 fiches de paire (statuts Figma)",
        "Liste 7:2229. Fiches : Active 7:2778, À graver 41:4277, Perdue 17:2118, Inactive 17:2009, À récupérer 17:2227, Livraison demandée 41:4164, En livraison 41:4022.",
        bullets(
            "Filtres type / statut.",
            "Titre type + marque, historique plus récent en premier.",
            "Menus d’actions selon le statut (mémos).",
        ),
        bullets(
            "Paire d’un autre magasin → inaccessible.",
            "Paire archivée : invisible magasin, visible admin.",
        ),
        "V1", "Non",
    ),
    (
        "FO-O-34", "Lunettes & codes", "US", "Modifier / déclarer inactive côté magasin — sans recycler le code",
        "Pop-in 7:3294. Excel LU-07 réécrit (Excel disait « libérer le code »). CDC §7.6 / §7.9.",
        bullets(
            "Modification libellé / type possible selon Figma.",
            "Inactive : cassée / perdue définitive, confirmation, motif, date, user.",
            "Code définitivement inutilisable.",
            "Blocage perte / découverte ultérieures. Arrêt des notifs de cycle.",
        ),
        bullets(
            "Overwrite du code assigné impossible.",
            "Inactive consultable en BO admin.",
        ),
        "V1", "Non",
    ),
    (
        "FO-O-40", "Restitution", "US", "Signaler une paire retrouvée (saisie du code assigné)",
        "Parcours 3 pop-ins : 7:3494, 7:4357, confirmations 7:3610… Variantes 41:4510+. CDC §7.7. Excel RE-02/RE-06.\nSaisie manuelle + champ compatible douchette (frappe clavier) CODE-01/09. Caméra CODE-07 = V1.1.",
        bullets(
            "Le système identifie la paire, crée le dossier, enregistre le dépositaire, génère le code de restitution, notifie le propriétaire, passe la paire À récupérer.",
            "Avertissement Figma : vérifier que c’est la bonne paire.",
            "Aucune identité propriétaire si ce n’est pas déjà un client de CE magasin.",
        ),
        bullets(
            "Code inconnu / mauvais format / mauvais statut → refus.",
            "Paire inactive / archivée → refus.",
            "Paire déjà À récupérer chez ce magasin → pas de 2e dossier ouvert (à figer : message).",
        ),
        "V1", "Email + SMS (notif client)",
    ),
    (
        "FO-O-41", "Restitution", "US", "Remise au client (code de restitution)",
        "Pop-ins 7:3540, erreur 41:3115, confirmer 7:3586. CDC §7.8. Excel RE-03/RE-04. Action définitive.",
        bullets(
            "Saisie / douchette du code de restitution.",
            "Nom+prénom seulement si client du magasin.",
            "Confirmation → paire Active, dossier sort de la liste, date + user tracés.",
            "CTA Remettre disabled s’il n’y a aucune paire À récupérer (FO-O-10).",
        ),
        bullets(
            "Mauvais code → pop-in erreur Figma.",
            "Abandon avant confirmer → toaster, dossier inchangé.",
        ),
        "V1", "Non",
    ),
    (
        "FO-O-42", "Restitution", "US", "Liste restitutions en cours + fiches dossier",
        "Liste 7:2444, fiches 7:2872 / 41:5776 / 41:5898. Excel RE-05 (Excel « tous dossiers » → Figma en cours only).",
        bullets(
            "Uniquement dossiers ouverts.",
            "Disparaît dès remise magasin OU livraison finalisée.",
            "Recherche client pour renvoi de code si oubli.",
        ),
        bullets(
            "Historique clôturé : pas dans ce tableau (visible admin BO-A-06).",
        ),
        "V1", "Non",
    ),
    (
        "FO-O-43", "Restitution", "US", "Renvoi du code de restitution (email et/ou SMS)",
        "Excel EC-03. CDC §8.4 : expiration / tentatives à figer (D-14).",
        bullets(
            "Renvoi selon consentements du client.",
            "Depuis fiche restitution / Relancer.",
            "Limitation de fréquence (anti-abus).",
        ),
        bullets(
            "Client sans tél / sans consentement SMS → email seul, et inversement.",
            "Dossier clôturé → pas de renvoi.",
        ),
        "V1", "Email + SMS",
    ),
    (
        "FO-O-44", "Envoi à distance", "US", "Traiter un envoi : étiquette, dépôt, bulk bordereaux, en livraison",
        "CDC traitement restitution à distance. Pop-in « Indiquer comme en livraison » 46:2177. Mémo bulk : dès 2 paires Livraison demandée, télécharger les bordereaux ; autres statuts non cochables.",
        bullets(
            "Notif après paiement confirmé.",
            "Étiquette prépayée imprimable et/ou QR (nom + adresse destinataire seulement).",
            "L’opticien ne paie pas l’affranchissement.",
            "Confirmer dépôt transporteur / relais → En livraison.",
            "Client informé du dépôt ; suivi si n° transporteur.",
            "Finalisation → paire Active, dossier hors liste en cours.",
        ),
        bullets(
            "Pas de paiement côté magasin.",
            "Étiquette avant paiement → interdit.",
            "ETA non affichée tant que contrat transporteur non figé.",
        ),
        "V1", "Colissimo / Mondial Relay (étiquette + suivi)",
    ),
    (
        "BO-A-01", "Admin plateforme", "US", "Connexion admin — espace distinct du magasin",
        "Excel AU-02. CDC §6 / §14. Back-office hors Figma.",
        bullets(
            "Authentification administrateur Perdues de Vue.",
            "Un opticien n’accède pas à cet espace (403).",
            "Un admin n’utilise pas le menu Figma magasin comme espace de travail.",
        ),
        bullets(
            "MDP oublié admin.",
            "Compte admin désactivé.",
        ),
        "V1", "Non",
    ),
    (
        "BO-A-02", "Admin plateforme", "US", "Dashboard KPIs globaux",
        "CDC §6.1 (liste ajustable) : réseaux, opticiens, users, clients, paires, codes dispo, codes associés, restitutions en cours / terminées, factures prêtes, factures en erreur, courbes 12 mois. Excel BO-01/BO-02.",
        bullets(
            "Indicateurs synthétiques consultables.",
            "Classement / activité opticiens (codes, restitutions).",
            "Compteurs factures alignés sur le contrôle de facturation.",
        ),
        bullets(
            "Plateforme vide (0 magasin).",
        ),
        "V1", "Non",
    ),
    (
        "BO-A-03", "Admin plateforme", "US", "Réseaux d’opticiens",
        "CDC §3.1 / §6.2 : raison sociale, description, prix par code, fréquence de facturation, actif, liste des magasins.",
        bullets(
            "Créer / modifier / consulter / activer-désactiver.",
            "Définir prix/code et fréquence.",
            "Consulter les opticiens du réseau.",
        ),
        bullets(
            "Désactiver un réseau avec magasins rattachés → règle à figer (blocage ou conservation héritage).",
        ),
        "V1", "Non",
    ),
    (
        "BO-A-04", "Admin plateforme", "US", "Opticiens : CRUD, blocage, facturation, héritage tarif",
        "CDC §3.2 / §3.3 / §6.3. Excel OP-05 / BO-03.\nFiche min. : RS, nom commercial, adresses exploitation et facturation, email, tél, SIRET, TVA intra si besoin, réseau 0 ou 1, prix et fréquence hérités ou surchargés, actif.",
        bullets(
            "Créer un magasin + envoi des accès.",
            "Rattacher / détacher d’un réseau.",
            "Héritage : (1) prix opticien si surcharge (2) sinon réseau (3) hors réseau : prix ET fréquence obligatoires.",
            "Activer / désactiver / bloquer.",
            "Consulter users et activité fonctionnelle.",
        ),
        bullets(
            "Opticien sans réseau et sans prix → facture en erreur (FACT-03).",
            "SIRET / adresse factu manquants → erreur de facturation, magasin métier possible (à figer si on bloque l’usage).",
        ),
        "V1", "Email invitation magasin",
    ),
    (
        "BO-A-05", "Admin plateforme", "US", "Utilisateurs plateforme et users magasin (rôles)",
        "CDC §6.4 / §7.2 / §2.2. Excel BO-06.\nRôles : admin_pdv, opticien_admin, opticien_user, client.\nUsers magasin : créer, modifier, activer, rôle, reset / invitation MDP — réservé admin magasin + admin PDV. Absent du menu Figma magasin.",
        bullets(
            "Recherche / filtres users.",
            "Rattacher un user à un seul opticien.",
            "opticien_user : métier sans gestion des comptes.",
            "opticien_admin : métier + users de l’établissement.",
        ),
        bullets(
            "Un admin magasin A ne gère pas les users de B.",
            "User désactivé → plus de login.",
        ),
        "V1", "Email invitation user magasin",
    ),
    (
        "BO-A-06", "Admin plateforme", "US", "Consultation transverse : clients, paires, codes, restitutions, notifications",
        "CDC §6.5 à §6.8. Excel BO-04/BO-05.",
        bullets(
            "Clients globaux : fiche, paires, opticien de rattachement, restitutions.",
            "Codes : recherche, filtre statut, lot d’origine, paire, opticien d’association, date, statut facturation.",
            "Restitutions en cours ET terminées : propriétaire, paire, dépositaire, code retrait, événements.",
            "Notifications : canal, destinataire, date, statut technique, dossier, erreurs.",
            "Paires inactives / archivées consultables ici.",
        ),
        bullets(
            "Pas d’action métier magasin qui contourne l’isolation (l’admin agit en conscience, journalisé).",
        ),
        "V1", "Non",
    ),
    (
        "BO-A-07", "Gravure", "US", "Lots de codes graveur (CODE-15 à 20) + 3 parcours CDC §10",
        "Parcours 1 : créer lot → volume → générer → export CSV → transmission graveur. Codes exportés réservés, jamais réutilisés, NON facturés tant que non associés.\nParcours 2 : saisie magasin code déjà gravé (FO-O-31).\nParcours 3 : génération à l’enregistrement (FO-O-30).\nDécision D-18 : confirmer si un magasin peut générer hors lot CSV.",
        bullets(
            "Historique des lots.",
            "Statut de chaque code du lot (généré, réservé/exporté, associé, facturé, désactivé, annulé technique).",
            "Annulation technique exceptionnelle : toujours sans remise à disposition du code.",
        ),
        bullets(
            "Export d’un lot vide.",
            "Volume trop élevé → plafond à figer.",
        ),
        "V1", "Non (CSV fichier)",
    ),
    (
        "BO-A-08", "Gravure", "US", "API / import automatique graveur (CODE-21 / 22)",
        "CDC 4.3 : alternative à l’export CSV manuel.",
        bullets(
            "Si retenu : transmission directe des lots au graveur.",
            "Sinon : rester sur CSV V1.",
        ),
        bullets(
            "Échec d’import graveur → lot reste réservé, pas recyclé.",
        ),
        "V1.1", "API graveur",
    ),
    (
        "FA-01", "Tarification", "US", "Prix unitaire par code associé — héritage réseau / surcharge magasin",
        "CDC §11.1. Unité : événement d’association code–paire. Réseaux non facturés.",
        bullets(
            "Prix au réseau ; magasin hérite ou surcharge.",
            "Hors réseau : prix propre obligatoire.",
            "Règle d’application (1)(2)(3) CDC.",
        ),
        bullets(
            "Changement de prix en cours de période : appliqué aux nouvelles associations après la date (règle de bascule à figer).",
            "Prix manquant → facture en erreur, pas d’émission.",
        ),
        "V1", "Non",
    ),
    (
        "FA-02", "Tarification", "US", "Tarification avancée TARIF-01 à TARIF-09",
        "CDC §11.2 : dégressifs volume, paliers, forfaits, abonnements, remises auto, promotions, tarif par type de code, par canal de génération, règles complexes.",
        bullets(
            "Chaque ID TARIF coché en atelier devient exigence de go-live ou de lot ultérieur.",
            "Préciser l’unité (paire / forfait / gravure séparée) — D-12.",
        ),
        bullets(
            "Combinaison de règles contradictoires → priorité commerciale à documenter.",
        ),
        "V1.1", "Non",
    ),
    (
        "FA-03", "Facturation", "US", "Fréquences et événement facturable",
        "CDC §12.1–12.2. Fréquences : quotidienne, hebdo, mensuelle, trimestrielle, semestrielle, annuelle. Héritage + surcharge.\nFacturable = nouvelles associations code–paire de la période. Lots générés / exportés / réservés non associés : NON facturés. Un événement n’est facturé qu’une fois.\nMontant HT = nb associations × prix unitaire applicable.",
        bullets(
            "Fréquence par défaut à figer (ex. mensuelle) — D-13.",
            "Préparation des volumes et du montant (FACT-01).",
        ),
        bullets(
            "Association puis inactive dans la même période : reste facturable (le code a été associé).",
            "Tentative de refacturer le même événement → interdit.",
        ),
        "V1", "Non",
    ),
    (
        "FA-04", "Facturation", "US", "Évaluation : factures prêtes vs en erreur + réévaluation (FACT-01 à 04)",
        "CDC §12.3 / §12.4. Erreurs types : SIRET absent, adresse factu incomplète, TVA, prix, fréquence, donnée requise par Qonto, incohérence technique.",
        bullets(
            "L’admin voit les factures prêtes (FACT-02) et en erreur (FACT-03) — aussi en KPI dashboard.",
            "Il corrige les données magasin / réseau puis relance (FACT-04).",
            "Après correction, passage possible à « prête à générer ».",
            "Une facture en erreur n’est pas transmise à Qonto.",
        ),
        bullets(
            "Correction partielle → toujours en erreur tant que le jeu n’est pas complet.",
            "Aucun opticien sur la période → 0 facture.",
        ),
        "V1", "Non",
    ),
    (
        "FA-05", "Facturation", "US", "Qonto : transmission, statuts, n° de facture, PDF (FACT-05 à 09)",
        "CDC §12.5 : Qonto est la solution de facturation PDV. L’admin évalue et transmet.",
        bullets(
            "Transmission en lot vers Qonto (FACT-05).",
            "Conservation du statut technique de transmission (FACT-06).",
            "Synchronisation des statuts de facture (FACT-07).",
            "Récupération du n° de facture (FACT-08) et du PDF (FACT-09).",
        ),
        bullets(
            "Échec de transmission → statut d’erreur, retry, pas de doublon d’émission.",
            "Facture déjà numérotée Qonto → pas de second envoi du même lot.",
        ),
        "V1", "Qonto",
    ),
    (
        "FA-06", "Facturation", "US", "Pilotage financier FACT-10 à 15",
        "Paiements, impayés, CA, tableaux financiers, rapprochement comptable, exports avancés.",
        bullets(
            "Exigences CDC 12.5 à prioriser en atelier (V1.1 proposé si le go-live se limite à l’émission Qonto).",
        ),
        bullets(
            "Dépend des FACT-07–09 (statuts / PDF Qonto).",
        ),
        "V1.1", "Qonto",
    ),
    (
        "FA-07", "Facturation", "US", "Facturation du service d’envoi — côté client (rappel catalogue)",
        "CDC §12.6. Voir aussi FO-C-23. Cycle indépendant des factures opticiens.",
        bullets(
            "Encaissement plateforme au moment de la demande.",
            "Justificatif client.",
            "Ne modifie pas le prix/code magasin.",
        ),
        bullets(
            "Remboursement / échec transporteur : règle financière à figer.",
        ),
        "V1", "Paiement en ligne",
    ),
    (
        "TR-01", "Codes", "US", "Saisie, format, unicité, copie (CODE-01 à 06, 12 à 14)",
        "Tous les champs code (paire, retrait) contrôlent format, existence, statut. Copie presse-papiers CODE-06.",
        bullets(
            "Alphanumérique, unicité globale.",
            "Format exact tranché en D-01.",
            "Code de restitution ≠ code assigné.",
        ),
        bullets(
            "Coller des espaces / casse → normalisation à figer (uppercase).",
            "Scanette HID : le champ se comporte comme une saisie clavier.",
        ),
        "V1", "Non",
    ),
    (
        "TR-02", "Codes", "US", "Caméra, QR auto, équipements tiers, reconnaissance (CODE-07 à 11)",
        "CDC 4.1 « toutes les modalités ». Figma = saisie + intention scanner. V1 proposé = saisie + douchette ; reste = V1.1.",
        bullets(
            "CODE-07 scan caméra smartphone/tablette.",
            "CODE-08 lecture auto QR / codes-barres.",
            "CODE-10 lecteurs tiers.",
            "CODE-11 reconnaissance au-delà du scan simple.",
        ),
        bullets(
            "Caméra refusée par l’OS → fallback saisie.",
        ),
        "V1.1", "Caméra / SDK lecture",
    ),
    (
        "TR-03", "Notifications", "US", "Emails / SMS transactionnels + journal + gabarit marque",
        "CDC §13. Excel NO-01/02/04/05. Cas : invitation, paire retrouvée + code retrait, localisation, renvoi code, confirmation restitution (si retenue), paiement envoi, notif opticien envoi payé, étiquette dispo, confirmation expédition, notifs techniques admin. Enrichis CDC = V1.1.",
        bullets(
            "Canaux email et SMS selon consentements.",
            "Journal : canal, destinataire, type, date, statut, erreur, dossier.",
            "Design Perdues de Vue (logo, couleurs, footer).",
            "Codes / OTP jamais loggés en clair dans les journaux métiers.",
        ),
        bullets(
            "Échec provider → statut erreur, pas de faux « envoyé ».",
            "Client sans consentement SMS → pas de SMS.",
        ),
        "V1", "Email + SMS",
    ),
    (
        "TR-04", "Notifications", "US", "Hors V1 proposé — notif perte à l’opticien d’origine (NO-03)",
        "Excel NO-03. Absent Figma et des cas CDC principaux. La découverte se fait par scan chez n’importe quel partenaire.",
        bullets(
            "Si hors V1 : aucune notif de déclaration de perte vers l’opticien d’origine.",
        ),
        bullets(
            "Si exigé en V1 : maquette et copy à créer.",
        ),
        "Hors V1", "Non",
    ),
    (
        "TR-05", "Sécurité", "US", "Isolation magasin, jetons publics, audit PII, MDP, 3 environnements",
        "CDC §14 / §15. Excel SE-03 / AR-01. Mémo PII dépositaire.",
        bullets(
            "Chaque requête magasin est limitée à son opticien.",
            "Client limité à ses paires.",
            "Jetons publics aléatoires, TTL, non séquentiels.",
            "Journalisation technique des accès PII et actions sensibles (alimente les écrans BO-A-09 / BO-A-10).",
            "MDP stockés de façon sécurisée.",
            "Dev / recette / prod séparés.",
            "Contrôles anti-recyclage des codes.",
        ),
        bullets(
            "IDOR : GET/PATCH/DELETE sur l’autre magasin / l’autre client → refus.",
            "Lien public d’un autre dossier → 404.",
        ),
        "V1", "Non",
    ),
    (
        "TR-06", "Sécurité", "US", "MFA pour rôles à privilèges (CDC §14)",
        "Exigence CDC, absente Figma. Décision D-16.",
        bullets(
            "Si V1 : au minimum admin_pdv.",
            "Sinon : lot sécurité ultérieur.",
        ),
        bullets(
            "Perte du second facteur → procédure support.",
        ),
        "V1.1", "Non",
    ),
    (
        "TR-07", "RGPD", "US", "Minimisation, information, conservation / effacement",
        "CDC §15. Non-communication des coordonnées au dépositaire sauf client préexistant ou consentement envoi. Politique à annexer (D-19).",
        bullets(
            "CGU / politique présentées à l’activation (SE-01).",
            "Accès limité selon les rôles.",
            "Anonymisation / désactivation selon règles retenues.",
        ),
        bullets(
            "Demande d’effacement avec restitution ouverte → procédure à figer.",
        ),
        "V1", "Non",
    ),
    (
        "BO-A-09", "Traçabilité BO", "US", "Journal des modifications — qui / quoi / quand / IP (front + BO)",
        "Interface admin dédiée pour consulter le tracé de toutes les actions métier, y compris celles réalisées depuis les front-offices client et opticien (et le BO).\nCDC §14 journalisation des opérations sensibles · Excel SE-03. Hors Figma (écran Filament).",
        bullets(
            "Chaque action significative produit une entrée : qui (utilisateur, rôle, magasin si applicable), quoi (action en français lisible), sur quoi (entité : client, paire, restitution, tarif…), quand (date/heure), depuis où (adresse IP).",
            "Texte lisible métier, pas un dump technique brut (ex. « Opticien Dupont a déclaré la paire FR5487OK retrouvée »).",
            "Couverture : création / modification / suppression / transitions de statut / restitutions / envois / facturation / users / paramètres — front client, front opticien et admin.",
            "Filtres : période (dates), utilisateur, rôle, magasin, type d’action, entité, IP.",
            "Tri par date décroissante ; pagination ; recherche texte sur le libellé.",
            "Consultation réservée admin_pdv ; pas d’édition / suppression des lignes (immuable).",
            "Pas de secrets en clair dans le journal (MDP, OTP, codes de restitution complets si politique de minimisation).",
        ),
        bullets(
            "Action anonyme via lien public restitution → entrée avec acteur « lien public » + IP.",
            "Échec d’action (refus policy) : optionnellement journalisé comme tentative (à figer).",
            "Volume élevé : filtres obligatoires / plage de dates max à figer.",
            "Export CSV des résultats filtrés (should).",
        ),
        "V1", "Non",
    ),
    (
        "BO-A-10", "Traçabilité BO", "US", "Journal des connexions — qui / quand / IP",
        "Interface admin séparée du journal des modifications, dédiée au tracé des authentifications.\nCDC §14. Hors Figma (écran Filament).",
        bullets(
            "Une entrée par événement de connexion : succès ou échec.",
            "Qui (email / identifiant tenté, user si connu, rôle), quand (date/heure), IP, espace (client / opticien / admin), résultat (succès / échec).",
            "Filtres : période, utilisateur, rôle, espace, succès/échec, IP.",
            "Tri par date décroissante ; pagination.",
            "Consultation réservée admin_pdv ; lignes immuables.",
            "Pas d’énumération facilitée côté login public ; le journal reste côté BO uniquement.",
        ),
        bullets(
            "Tentatives répétées (bruteforce) visibles via filtre échecs + IP.",
            "Déconnexion volontaire : optionnellement journalisée (à figer).",
            "Reset MDP / invitation consommée : plutôt dans le journal des modifications (BO-A-09), pas ici — sauf si on décide de les lister aussi en « événements d’accès ».",
        ),
        "V1", "Non",
    ),
]

# Ordre de lecture = parcours d’utilisation (client → opticien → admin/BO en dernier)
FLOW = [
    ("BANNER", 0, "1 · PARCOURS UTILISATEUR (CLIENT) — Figma page Client 1:535"),
    ("BANNER", 1, "Module A — Accéder au compte (invitation → connexion → reset)"),
    "FO-C-03",
    "FO-C-01",
    "FO-C-02",
    "FO-C-04",
    ("BANNER", 1, "Module B — Consulter Mes lunettes"),
    "FO-C-10",
    "FO-C-11",
    ("BANNER", 1, "Module C — Déclarer une perte / sortir du cycle"),
    "FO-C-12",
    "FO-C-13",
    "FO-C-14",
    ("BANNER", 1, "Module D — Paire retrouvée : localiser & récupérer en magasin"),
    "FO-C-05",
    "FO-C-15",
    ("BANNER", 1, "Module E — Se faire envoyer la paire (envoi à distance)"),
    "FO-C-20",
    "FO-C-21",
    "FO-C-22",
    "FO-C-23",
    "FO-C-16",
    "FO-C-24",
    ("BANNER", 1, "Module F — Paramètres du compte client"),
    "FO-C-30",
    "FO-C-31",
    ("BANNER", 0, "2 · PARCOURS OPTICIEN (MAGASIN) — Figma page Opticien 7:1919"),
    ("BANNER", 1, "Module A — Accéder au magasin"),
    "FO-O-01",
    "FO-O-02",
    ("BANNER", 1, "Module B — Tableau de bord & actions rapides"),
    "FO-O-10",
    ("BANNER", 1, "Module C — Gérer les clients"),
    "FO-O-20",
    "FO-O-21",
    "FO-O-22",
    ("BANNER", 1, "Module D — Enregistrer & suivre les paires (codes assignés)"),
    "FO-O-30",
    "FO-O-31",
    "FO-O-32",
    "FO-O-33",
    "FO-O-34",
    "TR-01",
    "TR-02",
    ("BANNER", 1, "Module E — Restituer en magasin (2 codes)"),
    "FO-O-40",
    "FO-O-41",
    "FO-O-42",
    "FO-O-43",
    ("BANNER", 1, "Module F — Traiter un envoi à distance"),
    "FO-O-44",
    ("BANNER", 1, "Module G — Notifications magasin"),
    "FO-O-12",
    "TR-04",
    ("BANNER", 1, "Module H — Paramètres compte opticien"),
    "FO-O-11",
    ("BANNER", 0, "3 · BACK-OFFICE ADMIN (en dernier) — CDC §6 / §11 / §12 — hors Figma"),
    ("BANNER", 1, "Module A — Accès admin"),
    "BO-A-01",
    ("BANNER", 1, "Module B — Pilotage (dashboard)"),
    "BO-A-02",
    ("BANNER", 1, "Module C — Réseaux & opticiens"),
    "BO-A-03",
    "BO-A-04",
    ("BANNER", 1, "Module D — Utilisateurs & rôles"),
    "BO-A-05",
    ("BANNER", 1, "Module E — Consultation transverse"),
    "BO-A-06",
    ("BANNER", 1, "Module F — Lots graveur"),
    "BO-A-07",
    "BO-A-08",
    ("BANNER", 1, "Module G — Tarification"),
    "FA-01",
    "FA-02",
    ("BANNER", 1, "Module H — Facturation opticiens & Qonto"),
    "FA-03",
    "FA-04",
    "FA-05",
    "FA-06",
    "FA-07",
    ("BANNER", 1, "Module I — Notifications système (journal / gabarits)"),
    "TR-03",
    ("BANNER", 1, "Module J — Sécurité & RGPD"),
    "TR-05",
    "TR-06",
    "TR-07",
    ("BANNER", 1, "Module K — Traçabilité (journaux BO)"),
    "BO-A-09",
    "BO-A-10",
]


def ordered_rows():
    stories = {item[0]: item for item in ROWS}
    out = []
    used = set()
    for step in FLOW:
        if isinstance(step, tuple) and step[0] == "BANNER":
            out.append(step)
            continue
        if step not in stories:
            raise SystemExit(f"FLOW référence un ID inconnu : {step}")
        out.append(stories[step])
        used.add(step)
    missing = set(stories) - used
    if missing:
        raise SystemExit(f"US absentes du FLOW : {sorted(missing)}")
    return out


def apply_border_fill(cell, fill=None, font=None, align=None):
    cell.border = thin
    cell.alignment = align or wrap
    if fill is not None:
        cell.fill = fill
    if font is not None:
        cell.font = font


def main():
    wb = Workbook()
    ws = wb.active
    ws.title = "Backlog"

    ws.merge_cells("A1:I1")
    t = ws["A1"]
    t.value = (
        "BACKLOG FINAL — Perdues de Vue  |  Une feuille, organisée par parcours d’usage  |  "
        "Ordre : 1 Client → 2 Opticien → 3 Back-office Admin (fin)  |  "
        "Sources : Figma > CDC > Excel 48  |  "
        + FIGMA
        + "  |  CDC §1–15.  "
        "Version : V1 · V1.1 · Hors V1.  "
        "Dépendance API = prestataire métier (pas un état d’avancement)."
    )
    t.font = Font(name="Calibri", bold=True, size=11, color=NAVY)
    t.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 48

    for i, h in enumerate(HEADERS, 1):
        cell = ws.cell(2, i, h)
        apply_border_fill(cell, fill_navy, font_h, Alignment(wrap_text=True, vertical="center", horizontal="center"))
    ws.row_dimensions[2].height = 28
    ws.freeze_panes = "A3"

    r = 3
    n_us = 0
    for item in ordered_rows():
        if item[0] == "BANNER":
            _, level, title = item
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=COLS)
            cell = ws.cell(r, 1, title)
            fill = fill_navy if level == 0 else fill_teal
            apply_border_fill(cell, fill, font_banner, wrap_c)
            for c in range(2, COLS + 1):
                apply_border_fill(ws.cell(r, c), fill)
            ws.row_dimensions[r].height = 22
            r += 1
            continue
        sid, epic, typ, resume, desc, ac, edges, ver, api = item
        vals = [sid, epic, typ, resume, desc, ac, edges, ver, api]
        fills = [None] * 9
        if ver == "V1":
            fills[7] = PatternFill("solid", fgColor="C6F0D4")
        elif ver == "V1.1":
            fills[7] = PatternFill("solid", fgColor="FDE8CD")
        elif ver == "Hors V1":
            fills[7] = PatternFill("solid", fgColor="F8D0D0")
        for c, val in enumerate(vals, 1):
            cell = ws.cell(r, c, val)
            fill = fills[c - 1]
            if fill is None and r % 2 == 0:
                fill = fill_grey
            apply_border_fill(cell, fill, font_b, wrap)
        ws.row_dimensions[r].height = 110
        n_us += 1
        r += 1

    widths = [12, 22, 8, 42, 62, 62, 52, 12, 28]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.auto_filter.ref = f"A2:I{r - 1}"
    dv = DataValidation(type="list", formula1='"V1,V1.1,Hors V1"', allow_blank=False)
    dv.add(f"H3:H{r - 1}")
    ws.add_data_validation(dv)

    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    wb.save(OUT)
    print("Wrote", OUT, "US", n_us, "rows", r - 1)


if __name__ == "__main__":
    main()
