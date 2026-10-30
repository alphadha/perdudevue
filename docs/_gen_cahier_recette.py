"""Génère docs/Cahier_de_recette_Perdues_de_Vue.xlsx — modèle Swiss Impact."""
from __future__ import annotations

from collections import Counter
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path(__file__).resolve().parent / "Cahier_de_recette_Perdues_de_Vue.xlsx"

HEADER_FILL = PatternFill("solid", fgColor="3658AC")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
TITLE_FONT = Font(name="Calibri", size=16, bold=True, color="1B2A4A")
SECTION_FONT = Font(name="Calibri", size=12, bold=True, color="3658AC")
BODY_FONT = Font(name="Calibri", size=11)
P0_FILL = PatternFill("solid", fgColor="FCE4E4")
P1_FILL = PatternFill("solid", fgColor="FFF4D6")
P2_FILL = PatternFill("solid", fgColor="E8F5E9")
ALT_FILL = PatternFill("solid", fgColor="F5F7FB")
THIN = Border(
    left=Side(style="thin", color="D0D5DD"),
    right=Side(style="thin", color="D0D5DD"),
    top=Side(style="thin", color="D0D5DD"),
    bottom=Side(style="thin", color="D0D5DD"),
)
WRAP = Alignment(wrap_text=True, vertical="top")

# (parcours, n, priorité, rôle, page, action, résultat, réf)
TESTS: list[tuple] = []


def t(parcours, n, prio, role, page, action, result, ref=""):
    TESTS.append((parcours, n, prio, role, page, action, result, ref))


# ─── P01 Fondations ───────────────────────────────────────────────────────────
P01 = "P01 — Fondations & smoke"
t(P01, 1, "P0", "QA", "API /api/v1/health", "Appeler GET health (local / recette)", "HTTP 200, payload santé OK.", "R0-01")
t(P01, 2, "P0", "Visiteur", "Front opticien (non connecté)", "Ouvrir une route opticien protégée sans session", "Redirect login. Aucune donnée métier visible.", "R0-02")
t(P01, 3, "P0", "Visiteur", "Front client (non connecté)", "Ouvrir /client (ou équivalent) sans session", "Redirect login client. Pas de liste de paires.", "R0-02")
t(P01, 4, "P0", "Visiteur", "Filament /admin", "Ouvrir /admin sans session", "Redirect login Filament. Pas de KPIs.", "R8-02")
t(P01, 5, "P1", "QA", "Front opticien", "Inspecter styles : classes pdv-* + Bootstrap 5", "Pas de feuille CSS custom hors tokens Prod UI / Bootstrap.", "R0-03")
t(P01, 6, "P0", "QA", "Env recette", "Vérifier sink mail (Mailpit/Mailtrap) + Stripe test + MR MODE", "Emails capturés ; Stripe test ; MR sandbox ou prod documenté.", "ENV")

# ─── P02 Auth opticien ────────────────────────────────────────────────────────
P02 = "P02 — Auth opticien"
t(P02, 1, "P0", "OPT-A", "/login opticien", "Se connecter avec email + MDP valides OPT-A", "Accès dashboard / home opticien.", "R1-01")
t(P02, 2, "P0", "OPT-A", "/login opticien", "Se déconnecter puis se reconnecter avec mauvais MDP", "Erreur générique (pas « email inconnu »). Pas de session.", "R1-02")
t(P02, 3, "P0", "OPT-A", "/login opticien", "Répéter N échecs login (throttle)", "Après N essais → 429 / message rate-limit. Pas d’énumération.", "R1-11")
t(P02, 4, "P0", "OPT-A", "Mot de passe oublié", "Demander reset avec email OPT-A", "Email sink avec code 6 chiffres ; pas de leak si email inconnu.", "R1-06")
t(P02, 5, "P0", "OPT-A", "Reset MDP", "Saisir le code + nouveau MDP conforme et valider", "MDP changé ; connexion OK avec le nouveau MDP.", "R1-06")
t(P02, 6, "P0", "OPT-A", "API admin", "Avec session OPT-A, appeler une route admin / Filament data", "403. Opticien n’accède pas au BO plateforme.", "R1-08")
t(P02, 7, "P1", "OPT-A user simple", "/login opticien", "Se connecter avec compte magasin non-admin (si existant)", "Accès magasin OK ; pas d’actions réservées admin magasin / BO.", "R2-10")
t(P02, 8, "P0", "OPT-A bloqué", "/login opticien", "Après blocage ADM (P10), tenter login OPT-A bloqué", "401 / erreur générique. Aucun accès.", "R8-04")

# ─── P03 Invitation client ────────────────────────────────────────────────────
P03 = "P03 — Invitation & auth client"
t(P03, 1, "P0", "C1", "Email invitation", "Ouvrir le lien d’invitation reçu (création client P04 ou seed)", "Page set password + CGU affichée. Token valide.", "R1-04")
t(P03, 2, "P0", "C1", "Set password", "Choisir un MDP sans cocher CGU et valider", "Bloqué. Compte non activé.", "R1-05")
t(P03, 3, "P0", "C1", "Set password", "Cocher CGU + MDP conforme et valider", "Compte activé. Accès espace client.", "R1-04")
t(P03, 4, "P0", "C1", "Lien invitation", "Réutiliser le même lien d’invitation", "Invalide / déjà utilisé. Pas de 2e activation.", "R1-07")
t(P03, 5, "P0", "C1", "/login client", "Tenter login avant set password (compte seed non activé)", "Refus. Message clair.", "R1-03")
t(P03, 6, "P0", "C1", "Reset client", "Mot de passe oublié → code email → nouveau MDP", "Reset OK ; connexion espace client.", "R1-06")
t(P03, 7, "P1", "C1", "Préférences", "1re connexion : observer prompt géoloc navigateur", "Prompt non bloquant si refus. App utilisable.", "R5-09")

# ─── P04 Clients magasin ──────────────────────────────────────────────────────
P04 = "P04 — Clients magasin"
t(P04, 1, "P0", "OPT-A", "Home / Clients", "Ouvrir la liste clients OPT-A", "Liste scoped magasin A. Pas de clients uniquement B.", "R2-06")
t(P04, 2, "P0", "OPT-A", "Nouveau client", "Créer C1 (email + tél + consentements email/SMS ON)", "Toaster succès ; fiche C1 ; email bienvenue en sink.", "R2-01")
t(P04, 3, "P0", "OPT-A", "Fiche C1", "Vérifier consentements email / SMS séparés", "Les deux persistés indépendamment.", "R2-02")
t(P04, 4, "P0", "OPT-A", "Nouveau client", "Recréer un client avec le même email que C1 → Confirmer rattachement", "Pas de 2e compte. C1 rattaché / flow doublon respecté.", "R2-03")
t(P04, 5, "P0", "OPT-A", "Doublon", "Sur doublon, choisir « Mettre à jour » (OTP selon canal)", "OTP reçu ; après validation, retour flow OK.", "R2-04")
t(P04, 6, "P1", "OPT-A", "Doublon", "Sur doublon, cliquer Annuler", "Formulaire intact. Aucune création.", "R2-05")
t(P04, 7, "P0", "OPT-A", "Fiche C1", "Changer l’email → OTP mail", "Pas de persist email avant OTP OK.", "R2-08")
t(P04, 8, "P0", "OPT-A", "Fiche C1", "Changer le téléphone → OTP SMS (sandbox)", "OTP SMS ; tél mis à jour après validation.", "R2-09")
t(P04, 9, "P0", "OPT-A", "Clients", "Unlink / détacher C1 chez A alors que C1 aussi chez B (préparer C2 ou rattachement B)", "B conserve le client. Isolation unlink.", "R2-07")
t(P04, 10, "P0", "OPT-A", "Navigation", "Chercher un menu « Utilisateurs magasin » dans Next", "Absent. Comptes magasin = Filament fiche opticien uniquement.", "R2-10")
t(P04, 11, "P0", "ADM", "Filament opticien", "Inviter un user magasin depuis fiche opticien Filament", "Email invitation ; autre magasin ne voit pas cet user (404 service).", "R2-11")
t(P04, 12, "P0", "OPT-B", "Clients", "Liste clients OPT-B : C1 uniquement A n’apparaît pas", "Isolation tenancy. C1 invisible chez B.", "R2-06")

# ─── P05 Paires & codes ───────────────────────────────────────────────────────
P05 = "P05 — Paires & codes"
t(P05, 1, "P0", "OPT-A", "Enregistrer une paire", "Créer P2 sans code (client C1, type Vue, marque)", "Statut À graver ; code auto format FR + 4 chiffres + 2 lettres.", "R3-01")
t(P05, 2, "P1", "OPT-A", "Modal paire", "Tester autocomplete client + marque + « Autre marque »", "Libellé libre marque persisté.", "R3-02")
t(P05, 3, "P1", "OPT-A", "Modal paire", "Choisir type Vue / Soleil / Solaire de Vue uniquement", "Seuls ces 3 types V1 proposés.", "R3-03")
t(P05, 4, "P2", "OPT-A", "Modal paire", "Sélectionner plusieurs couleurs et enregistrer", "Couleurs multi-select persistées.", "R3-04")
t(P05, 5, "P1", "OPT-A", "Fiche client → Nouvelle paire", "Ouvrir création paire depuis fiche C1", "Client prérempli non éditable (MX 7:3480).", "R3-05")
t(P05, 6, "P0", "OPT-A", "Fiche P2 À graver", "Exporter le code à graver", "Code visible / exportable. Format FR5487OK-like.", "R3-01")
t(P05, 7, "P0", "OPT-A", "Fiche P2", "Indiquer comme gravée (confirmation)", "→ Active. Action irréversible.", "R3-06")
t(P05, 8, "P0", "OPT-A", "Enregistrer paire", "Créer paire avec code déjà lié à une autre", "422. Pas d’overwrite. Message clair.", "R3-07")
t(P05, 9, "P0", "OPT-A", "Paire Active", "Déclarer inactive (confirmation)", "Inactive. Plus de « retrouvée » ni « perdue » possibles.", "R3-08")
t(P05, 10, "P0", "OPT-A / ADM", "Archive", "Archiver / supprimer une paire (API ou BO) puis retenter le même code", "Code non réutilisable.", "R3-09")
t(P05, 11, "P2", "OPT-A", "Fiche paire", "Vérifier titre = type + marque", "Conforme MX 32:2014.", "R3-10")
t(P05, 12, "P1", "OPT-A", "Fiche paire · Historique", "Créer plusieurs events puis ouvrir l’historique", "Plus récent en haut (MX 32:3752).", "R3-11")
t(P05, 13, "P0", "OPT-A", "Liste / fiche À graver", "Vérifier CTAs : exporter + indiquer gravée ; Remettre/Relancer/bordereau disabled", "Menus ⋮ conformes Figma lot 3.", "FIG-L3")
t(P05, 14, "P0", "OPT-A", "Seed P1", "S’assurer d’une paire Active P1 code FR5487OK (ou équivalent seed) pour P06", "P1 Active prête pour cycle restitution.", "DATA")

# ─── P06 Restitution critique ─────────────────────────────────────────────────
P06 = "P06 — Restitution magasin (critique)"
t(P06, 1, "P0", "C1", "Espace client", "Sur P1 Active, déclarer perdue + confirmer", "Statut Perdue. Event créé. Pas de notif magasin d’origine (NO-03).", "R5-02 · R6-06")
t(P06, 2, "P0", "OPT-B", "Signaler retrouvé", "Saisir le code assigné P1 (C1 n’est PAS client de B)", "Dossier restitution créé. Aucun nom/email/tél C1 affiché.", "R4-01")
t(P06, 3, "P0", "C1", "Email sink", "Ouvrir l’email « paire retrouvée »", "Lien public token + code restitution + magasin B. Pas de PII tiers.", "R4-02 · R6-03")
t(P06, 4, "P0", "Visiteur", "/r/{token}", "Ouvrir le lien public sans session", "Page visible ; une seule paire ; infos magasin B.", "R1-09")
t(P06, 5, "P0", "Visiteur", "/r/{token}", "Modifier le token (tamper) et recharger", "404. Pas de dump.", "R1-10")
t(P06, 6, "P0", "OPT-B", "Home opticien", "Observer le bandeau / liste d’attente (à récupérer)", "P1 visible. Max 5 sur dashboard (MX 41:2890).", "R4-04")
t(P06, 7, "P0", "OPT-B", "CTA Remettre", "Si 0 paire à récupérer (autre magasin / empty)", "CTA Remettre disabled.", "R4-03")
t(P06, 8, "P0", "OPT-B", "Remettre P1", "Lancer remise : mauvais code restitution", "Erreur Figma. Pas de clôture. Statut inchangé.", "R4-07")
t(P06, 9, "P0", "OPT-B", "Remettre P1", "Code restitution OK (C1 inconnu de B) → confirmer", "Remise OK. Identité masquée à l’étape 2. Paire → Active côté client.", "R4-05")
t(P06, 10, "P0", "OPT-B", "Liste restitutions", "Après remise, vérifier liste", "P1 absente des restitutions en cours.", "R4-09")
t(P06, 11, "P0", "OPT-A", "Remise C2", "Préparer C2 client A+B ; paire à récupérer chez A ; remettre", "Nom prénom visibles (déjà client local).", "R4-06")
t(P06, 12, "P0", "OPT-B", "Signaler retrouvé", "Scanner / saisir une paire Inactive", "Refus. Pas de dossier.", "R4-08")
t(P06, 13, "P1", "C1 / OPT-B", "Code oublié", "Demander renvoi code restitution (mail puis SMS)", "Renvoi OK + rate-limit après N essais.", "R4-10")
t(P06, 14, "P0", "OPT-B", "Inbox", "À la confirmation « retrouvée », vérifier notification inbox", "Titre « Lunettes en attente de retrait » ; nom seulement si déjà client.", "R6-04")

# ─── P07 Espace client ────────────────────────────────────────────────────────
P07 = "P07 — Espace client"
t(P07, 1, "P0", "C1", "Mes lunettes", "Ouvrir la liste des paires", "Ordre : à récupérer → en livraison → livraison demandée → perdue → active → inactive.", "R5-01")
t(P07, 2, "P0", "C1", "Fiche Active", "Déclarer perdue + confirmer", "→ Perdue. Event. CTA conformes.", "R5-02")
t(P07, 3, "P1", "C1", "Fiche Perdue", "Annuler la perte (si proposé)", "Retour Active.", "R5-03")
t(P07, 4, "P0", "C1", "Fiche Active", "Déclarer inactive + confirmer", "Irréversible. Plus de déclaration perte. « Réactiver » visible mais disabled.", "R5-04")
t(P07, 5, "P0", "C1", "Fiche", "Supprimer définitivement une paire (parcours prévu)", "Disparaît client + opticien ; visible BO taguée archivée.", "R5-05")
t(P07, 6, "P0", "C1", "À récupérer", "Ouvrir page / section À récupérer (après P06)", "Code restitution + Maps magasin. Pas de PII d’autres clients.", "R5-06")
t(P07, 7, "P0", "C1", "Liste", "Vérifier isolation : C1 ne voit aucune paire de C2", "Isolation stricte.", "R5-07")
t(P07, 8, "P1", "C1", "Préférences notif", "Couper SMS puis déclencher un événement SMS", "Aucun SMS envoyé (log skipped / absent).", "R5-08")
t(P07, 9, "P0", "C1", "À récupérer", "CTA « Me faire livrer » si a_recuperer", "CTA actif (entrée lot 7 / P09).", "R7")

# ─── P08 Notifications ────────────────────────────────────────────────────────
P08 = "P08 — Notifications & inbox"
t(P08, 1, "P0", "QA / ADM", "Journal notif", "Pour invitation, OTP, paire retrouvée, renvoi SMS : vérifier log", "Chaque envoi : canal, destinataire, type, statut.", "R6-01")
t(P08, 2, "P1", "QA", "Provider KO", "Simuler échec provider (clé invalide / mock)", "Statut error en log ; pas de 500 exposé au front.", "R6-02")
t(P08, 3, "P0", "OPT-B", "Inbox magasin", "Ouvrir inbox après signalement retrouvé", "Notification présente ; scoped magasin.", "R6-04")
t(P08, 4, "P0", "OPT-A", "Inbox", "OPT-A tente d’accéder inbox / notif d’OPT-B (IDOR)", "Liste vide / 404. Pas de fuite.", "R6-05")
t(P08, 5, "P0", "OPT-B", "Inbox", "Tout marquer comme lu puis Tout supprimer (onglets Toutes / Non lues / Lues)", "Actions scoped magasin. Onglets OK.", "R6-07")
t(P08, 6, "P0", "C1", "Email", "Lien email restitution = token public", "URL /r/{token} fonctionne (PDV-104).", "R6-03")

# ─── P09 Envoi à distance ─────────────────────────────────────────────────────
P09 = "P09 — Envoi à distance"
t(P09, 1, "P0", "C1", "À récupérer", "Démarrer « Me faire livrer » sans cocher consentement adresse", "422 consent_address. Aucun shipment créé.", "R7-02")
t(P09, 2, "P0", "C1", "Parcours livraison", "Choisir mode (relais / domicile) → récap → Procéder au paiement", "Checkout Stripe pending_payment ; paire reste a_recuperer jusqu’au retour OK.", "R7-01")
t(P09, 3, "P0", "C1", "Stripe test", "Payer avec 4242…4242 (succès)", "Retour succès ; toast succès ; statut livraison_demandee ; inbox dépositaire.", "R7-01 · R7-03")
t(P09, 4, "P0", "C1", "Stripe test", "Refaire un parcours et payer avec 4000…0002 (échec)", "Retour récap ?paiement=echec ; toast échec ; pas d’inbox ; pas livraison_demandee.", "R7-07")
t(P09, 5, "P0", "OPT-B", "Inbox / bordereau", "Ouvrir notification bordereau après paiement OK", "Nom + adresse uniquement (pas email/tél). OPT-A ne peut pas télécharger (404).", "R7-03")
t(P09, 6, "P0", "OPT-B", "Télécharger bordereau", "Télécharger le label (MR sandbox ou prod selon env)", "PDF si credentials OK ; sinon .txt fallback documenté. Pas de crash.", "MR")
t(P09, 7, "P0", "OPT-B", "Liste lunettes", "Multi-sélection : seules « Livraison demandée » cochables ; CTA dès 2", "Autres checkboxes disabled (MX 41:5612).", "R7-04")
t(P09, 8, "P0", "OPT-B", "Fiche / action", "Indiquer en livraison", "→ en_livraison ; mail client ; event « Livraison prise en compte. »", "R7-05")
t(P09, 9, "P1", "C1 / OPT-B", "Suivi", "Vérifier absence d’ETA tant que contrat transporteur incomplet", "Pas de champ eta API ni UI (MX 63:2111).", "R7-06")
t(P09, 10, "P0", "OPT-B", "Dépôt", "Saisir / confirmer dépôt avec code suivi", "Mail dépôt ; suivi visible côté client « Suivre ma livraison » ; pas d’ETA fantôme.", "R7-08")
t(P09, 11, "P0", "C1", "Espace client", "Ouvrir suivi livraison", "Statut / tracking visibles selon data. Isolation C1 only.", "R7-08")

# ─── P10 Admin Filament ───────────────────────────────────────────────────────
P10 = "P10 — Admin Filament plateforme"
t(P10, 1, "P0", "ADM", "/admin login", "Se connecter admin_pdv", "Dashboard Filament accessible.", "R8-02")
t(P10, 2, "P0", "OPT-A", "/admin", "Tenter /admin avec compte opticien", "403 / redirect. Pas de KPIs.", "R8-02")
t(P10, 3, "P0", "ADM", "Dashboard", "Vérifier KPIs + consultation paires (y compris archivées) + restitutions en cours + clients", "Données présentes. Restitutions = en cours only.", "R8-02")
t(P10, 4, "P0", "ADM", "Journal notifications", "Ouvrir journal notif BO", "Canal / destinataire / type / statut visibles.", "R8-01")
t(P10, 5, "P0", "ADM", "Opticiens", "Bloquer un compte opticien test", "Login opticien → 401 générique ensuite.", "R8-04")
t(P10, 6, "P1", "ADM", "Réseaux / opticiens", "CRUD réseau + rattacher opticien ; tester SIRET doublon", "Doublon SIRET refusé. Opticien sans SIRET/réseau listé au contrôle.", "R10-01")
t(P10, 7, "P1", "ADM", "Brand / UI", "Vérifier logo PDV + couleurs / CTA sombre Filament", "Conforme charte BO (hors Figma).", "UI-BO")

# ─── P11 Lots de codes ────────────────────────────────────────────────────────
P11 = "P11 — Lots de codes (gravure admin)"
t(P11, 1, "P0", "ADM", "Lots de codes", "Générer un lot de N codes", "Codes uniques FR####AA ; export CSV ; jamais recyclés.", "R9-01")
t(P11, 2, "P0", "OPT-A", "Enregistrer paire", "Saisir un code du lot comme « déjà gravé »", "Item lot → assigned ; paire Active.", "R9-02")
t(P11, 3, "P0", "OPT-A", "Enregistrer paire", "Retenter le même code du lot", "422. Pas de réutilisation.", "R9-02")
t(P11, 4, "P0", "ADM", "Lot", "Consommer tous les codes dispos → statut lot", "Lot consumed si plus aucun dispo.", "R9-02")
t(P11, 5, "P0", "OPT-A", "BO lots", "Opticien tente d’ouvrir consultation codes lot BO", "403.", "R9-03")
t(P11, 6, "P0", "ADM", "Consultation BO", "admin_pdv voit les codes du lot", "Liste complète + statuts.", "R9-03")

# ─── P12 Facturation ──────────────────────────────────────────────────────────
P12 = "P12 — Facturation & Qonto"
t(P12, 1, "P0", "ADM", "Tarifs", "Créer / éditer un tarif réseau (à la paire)", "CRUD OK. Montant indicatif cohérent.", "R10-01")
t(P12, 2, "P1", "ADM", "Préparation mois", "Lancer préparation mois de facturation", "Paires du mois + montant indicatif si tarif à la paire ; pas de sync Qonto auto.", "R10-04")
t(P12, 3, "P0", "ADM", "Qonto sync", "Sans clé Qonto : tenter sync", "Sync refusée ; log skipped. Pas d’appel HTTP.", "R10-02")
t(P12, 4, "P1", "ADM", "Qonto sync", "Avec clé : sync → statut queued sans appel HTTP réel (V1)", "queued documenté. Pas d’erreur 500.", "R10-02")
t(P12, 5, "P0", "ADM", "Billing review", "Ouvrir /admin/billing-review et /admin/qonto-sync", "Accessible admin_pdv.", "R10-03")
t(P12, 6, "P0", "OPT-A", "Billing BO", "Opticien tente billing-review / qonto-sync", "403.", "R10-03")
t(P12, 7, "P1", "ADM", "Tables billing Filament", "Vérifier lisibilité tables issues / previews", "Colonnes claires ; pas de PII inutile.", "UI-BO")

# ─── P13 Sécurité transverse ──────────────────────────────────────────────────
P13 = "P13 — Sécurité transverse"
t(P13, 1, "P0", "C1", "API paire", "GET paire appartenant à C2", "404 (IDOR).", "S-01")
t(P13, 2, "P0", "OPT-A", "API client", "GET client uniquement chez B", "404.", "S-02")
t(P13, 3, "P0", "OPT-A", "API PATCH profil", "Tenter mass-assignment role=admin", "Rôle inchangé.", "S-03")
t(P13, 4, "P0", "OPT-A", "Libellé paire", "Saisir <script>alert(1)</script> dans marque / libellé", "Échappé. Pas d’exécution XSS.", "S-04")
t(P13, 5, "P0", "OPT-A", "Recherche", "Injecter ' OR 1=1 -- dans recherche clients", "Ignorée / paramétrée. Pas de dump.", "S-05")
t(P13, 6, "P0", "QA", "Audit", "Après une remise, vérifier entrée audit", "Audit présent (remise + accès fiche sensibles).", "S-06")
t(P13, 7, "P0", "QA", "DB users", "Vérifier MDP hashés (bcrypt/argon) — pas de clair", "Hash only.", "S-07")
t(P13, 8, "P0", "QA", "CORS", "Origine non autorisée → API", "CORS refuse hors front recette.", "S-08")
t(P13, 9, "P0", "QA", "Logs", "Déclencher OTP / code restitution / login", "Pas de codes / MDP / OTP en clair dans laravel.log.", "SEC")
t(P13, 10, "P0", "QA", "Tokens publics", "Vérifier tokens restitution / invitation", "Aléatoires, TTL, non séquentiels.", "SEC")

# ─── P14 Parcours bout-en-bout magasin (référence) ─────────────────────────────
P14 = "P14 — Parcours magasin de référence (E2E)"
t(P14, 1, "P0", "ADM", "Filament", "Créer / vérifier opticien A et B + users", "Deux magasins opérationnels.", "PARCOURS")
t(P14, 2, "P0", "OPT-A", "Clients", "Créer client C1 chez A", "C1 OK + invitation.", "PARCOURS")
t(P14, 3, "P0", "C1", "Invitation", "Activer compte (MDP + CGU)", "Espace client OK.", "PARCOURS")
t(P14, 4, "P0", "OPT-A", "Paires", "Créer paire (avec ou sans gravure) → Active", "Code assigné unique FR…", "PARCOURS")
t(P14, 5, "P0", "C1", "Client", "Déclarer la paire perdue", "Perdue.", "PARCOURS")
t(P14, 6, "P0", "OPT-B", "Retrouvé", "Signaler retrouvée avec code assigné", "Dossier sans PII propriétaire.", "PARCOURS")
t(P14, 7, "P0", "C1", "Email + /r/token", "Récupérer code restitution", "Code + magasin B.", "PARCOURS")
t(P14, 8, "P0", "OPT-B", "Remettre", "Valider code restitution → confirmer remise", "Remise OK ; identité masquée si non client B.", "PARCOURS")
t(P14, 9, "P0", "C1", "Mes lunettes", "Vérifier paire Active à nouveau", "Active. Code assigné non réutilisable ailleurs.", "PARCOURS")
t(P14, 10, "P0", "QA", "GO", "Aucun cas P0 KO sur P01–P13 bloquant", "GO livraison possible si P0 = 0 ouvert.", "GO")


PARCOURS_META = [
    (1, P01, "QA / Visiteur", "Env, health, guards, config transporteurs"),
    (2, P02, "OPT-A", "Login, reset, throttle, isolation rôles"),
    (3, P03, "C1", "Invitation, CGU, set password, reset client"),
    (4, P04, "OPT-A / OPT-B / ADM", "CRUD clients, doublons, OTP, tenancy, users BO"),
    (5, P05, "OPT-A", "Création paires, gravure, inactive, immuabilité codes"),
    (6, P06, "OPT-B / C1", "Chemin critique trouvé → remise + PII"),
    (7, P07, "C1", "Liste, perte, inactive, à récupérer, prefs"),
    (8, P08, "OPT / ADM / C1", "Journal, inbox, IDOR notif, liens token"),
    (9, P09, "C1 / OPT-B", "Stripe, bordereau, multi-select, suivi"),
    (10, P10, "ADM", "Filament KPIs, journal, blocage, réseaux"),
    (11, P11, "ADM / OPT-A", "Lots CSV, assignation, non-réutilisation"),
    (12, P12, "ADM", "Tarifs, mois, Qonto, accès billing"),
    (13, P13, "QA", "IDOR, XSS, SQLi, audit, secrets, CORS"),
    (14, P14, "Tous", "E2E magasin de référence (DoD lotissement)"),
]


def style_header(ws, ncols: int):
    for c in range(1, ncols + 1):
        cell = ws.cell(1, c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = THIN


def autosize(ws, widths: dict[str, float]):
    for col, w in widths.items():
        ws.column_dimensions[col].width = w


def write_mode_emploi(wb: Workbook):
    ws = wb.create_sheet("Mode d'emploi", 0)
    ws["A1"] = "Cahier de recette — Perdues de Vue V1"
    ws["A1"].font = TITLE_FONT
    ws.merge_cells("A1:B1")

    rows = [
        (3, "Principe de test", None),
        (
            4,
            "Parcours continus",
            "Chaque parcours se joue dans l’ordre N°1 → N°2 → N°3… Ne pas sauter d’étape : chaque action part de l’écran atteint à l’étape précédente.",
        ),
        (
            5,
            "Rôles l’un après l’autre",
            "Ordre recommandé : Fondations → Auth opticien → Invitation client → Clients → Paires → Restitution (critique) → Espace client → Notifications → Envoi à distance → Admin → Lots codes → Facturation → Sécurité → E2E magasin.",
        ),
        (
            6,
            "Respect de l’app",
            "Suivre les routes / écrans réels (opticien Next, client Next, Filament /admin, page publique /r/{token}). Utiliser la navigation de l’app, pas des deep-links fantaisistes.",
        ),
        (8, "Colonnes (feuille Tests)", None),
        (9, "Parcours", "Nom du scénario fluide (P01, P02…)."),
        (10, "N°", "Ordre strict dans le parcours."),
        (11, "Priorité", "P0 bloquant · P1 important métier · P2 confort / non-régression."),
        (12, "Rôle", "Qui exécute l’étape (OPT-A, OPT-B, C1, ADM, QA…)."),
        (13, "Page / écran", "Où tu es (ou dois arriver) pour cette étape."),
        (14, "Action", "Ce que tu fais maintenant — suite logique de l’étape d’avant."),
        (15, "Résultat attendu", "Ce qui doit être vrai pour cocher OK."),
        (16, "Réf.", "ID cahier md (R4-01…), mémo Figma, ticket."),
        (17, "Statut", "OK / KO / N/A / Reporté / À retester."),
        (18, "Commentaire testeur", "Preuve (capture, curl, n° mail sink) — obligatoire si KO ou cas P0."),
        (20, "Priorités", None),
        (21, "P0", "Bloquant livraison / GO lot. IDOR, PII, code réutilisé, remise sans code = NOGO."),
        (22, "P1", "Important métier. Plan d’action daté si ouvert."),
        (23, "P2", "Confort / non-régression visuelle."),
        (25, "Jeux de données", None),
        (26, "ADM", "Admin PDV (Filament admin_pdv)."),
        (27, "OPT-A", "Opticien A — magasin d’origine de C1 / P1."),
        (28, "OPT-B", "Opticien B — dépositaire (trouve la paire de C1)."),
        (29, "C1", "Client rattaché à A uniquement (email + tél, consentements ON)."),
        (30, "C2", "Client rattaché à A et B (affichage nom à la remise)."),
        (31, "P1", "Paire C1, code assigné type FR5487OK — cycle complet."),
        (32, "P2", "Paire À graver (création sans code)."),
        (33, "DUP", "Tentative doublon email de C1."),
        (35, "Environnements", None),
        (36, "local", "Dev — Mailpit, Stripe test, MR selon .env."),
        (37, "recette", "Recette client / PO."),
        (38, "prod", "Réel — ne pas exécuter les cas destructifs sans plan."),
        (40, "GO / NOGO", None),
        (
            41,
            "GO lot",
            "0 ticket P0/C ouvert. Les P1/M ont un plan. Un IDOR ou un code réutilisé = NOGO immédiat.",
        ),
        (
            42,
            "Sources",
            "docs/04-cahier-de-recette.md · 01-regles-ux-memos-figma.md · 02-lotissement.md · 03-backlog-technique.md · Figma Prod UI BAT.",
        ),
    ]
    for r, a, b in rows:
        ws.cell(r, 1, a).font = SECTION_FONT if b is None else Font(name="Calibri", size=11, bold=True)
        if b is not None:
            ws.cell(r, 2, b).font = BODY_FONT
            ws.cell(r, 2).alignment = WRAP
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 100


def write_parcours(wb: Workbook):
    ws = wb.create_sheet("Parcours", 1)
    headers = ["Ordre", "Parcours", "Rôle principal", "Nb étapes", "Objectif"]
    for i, h in enumerate(headers, 1):
        ws.cell(1, i, h)
    style_header(ws, 5)

    counts = Counter(row[0] for row in TESTS)
    for ordre, name, role, objectif in PARCOURS_META:
        r = ordre + 1
        ws.cell(r, 1, ordre).font = BODY_FONT
        ws.cell(r, 2, name).font = BODY_FONT
        ws.cell(r, 3, role).font = BODY_FONT
        ws.cell(r, 4, counts[name]).font = BODY_FONT
        ws.cell(r, 5, objectif).font = BODY_FONT
        for c in range(1, 6):
            ws.cell(r, c).border = THIN
            ws.cell(r, c).alignment = WRAP
            if ordre % 2 == 0:
                ws.cell(r, c).fill = ALT_FILL
    autosize(ws, {"A": 8, "B": 48, "C": 22, "D": 12, "E": 58})
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:E{len(PARCOURS_META) + 1}"


def write_tests(wb: Workbook):
    ws = wb.create_sheet("Tests", 2)
    headers = [
        "Parcours",
        "N°",
        "Priorité",
        "Rôle",
        "Page / écran",
        "Action (suite de l’étape précédente)",
        "Résultat attendu",
        "Réf.",
        "Statut",
        "Commentaire testeur",
    ]
    for i, h in enumerate(headers, 1):
        ws.cell(1, i, h)
    style_header(ws, 10)
    ws.row_dimensions[1].height = 28

    prio_fill = {"P0": P0_FILL, "P1": P1_FILL, "P2": P2_FILL}
    for i, row in enumerate(TESTS, 2):
        parcours, n, prio, role, page, action, result, ref = row
        values = [parcours, n, prio, role, page, action, result, ref, None, None]
        for c, v in enumerate(values, 1):
            cell = ws.cell(i, c, v)
            cell.font = BODY_FONT
            cell.border = THIN
            cell.alignment = WRAP
        ws.cell(i, 3).fill = prio_fill.get(prio, PatternFill())
        if i % 2 == 0:
            for c in (1, 2, 4, 5, 6, 7, 8, 9, 10):
                if ws.cell(i, c).fill.fgColor is None or ws.cell(i, c).fill.fgColor.rgb in (
                    "00000000",
                    "0",
                ):
                    ws.cell(i, c).fill = ALT_FILL

    autosize(
        ws,
        {
            "A": 42,
            "B": 6,
            "C": 10,
            "D": 16,
            "E": 34,
            "F": 56,
            "G": 56,
            "H": 16,
            "I": 12,
            "J": 28,
        },
    )
    ws.freeze_panes = "A2"
    last = len(TESTS) + 1
    ws.auto_filter.ref = f"A1:J{last}"

    dv = DataValidation(
        type="list",
        formula1='"OK,KO,N/A,Reporté,À retester"',
        allow_blank=True,
    )
    dv.error = "Choisir une valeur de la liste"
    dv.errorTitle = "Statut"
    ws.add_data_validation(dv)
    dv.add(f"I2:I{last}")

    for r in range(2, last + 1):
        ws.row_dimensions[r].height = 48


def write_synthese(wb: Workbook):
    ws = wb.create_sheet("Synthèse", 3)
    headers = ["Parcours", "Nb étapes", "P0", "P1", "P2"]
    for i, h in enumerate(headers, 1):
        ws.cell(1, i, h)
    style_header(ws, 5)

    by = {}
    for parcours, _n, prio, *_rest in TESTS:
        by.setdefault(parcours, Counter())[prio] += 1
        by[parcours]["total"] += 1

    r = 2
    tot = Counter()
    for _ordre, name, *_ in PARCOURS_META:
        c = by[name]
        ws.cell(r, 1, name).font = BODY_FONT
        ws.cell(r, 2, c["total"]).font = BODY_FONT
        ws.cell(r, 3, c["P0"]).font = BODY_FONT
        ws.cell(r, 4, c["P1"]).font = BODY_FONT
        ws.cell(r, 5, c["P2"]).font = BODY_FONT
        ws.cell(r, 3).fill = P0_FILL
        ws.cell(r, 4).fill = P1_FILL
        ws.cell(r, 5).fill = P2_FILL
        for col in range(1, 6):
            ws.cell(r, col).border = THIN
        tot["total"] += c["total"]
        tot["P0"] += c["P0"]
        tot["P1"] += c["P1"]
        tot["P2"] += c["P2"]
        r += 1

    ws.cell(r, 1, "TOTAL").font = Font(name="Calibri", size=11, bold=True)
    ws.cell(r, 2, tot["total"]).font = Font(name="Calibri", size=11, bold=True)
    ws.cell(r, 3, tot["P0"]).font = Font(name="Calibri", size=11, bold=True)
    ws.cell(r, 4, tot["P1"]).font = Font(name="Calibri", size=11, bold=True)
    ws.cell(r, 5, tot["P2"]).font = Font(name="Calibri", size=11, bold=True)
    for col in range(1, 6):
        ws.cell(r, col).border = THIN
        ws.cell(r, col).fill = PatternFill("solid", fgColor="D6E0F5")

    r += 2
    ws.cell(r, 1, "Règle GO").font = SECTION_FONT
    ws.cell(r + 1, 1, "GO si 0 P0 ouvert et plan daté sur les P1. NOGO si 1 P0 (PII, IDOR, code réutilisé, remise sans code).").font = BODY_FONT
    ws.merge_cells(start_row=r + 1, start_column=1, end_row=r + 1, end_column=5)

    autosize(ws, {"A": 48, "B": 12, "C": 8, "D": 8, "E": 8})
    ws.freeze_panes = "A2"


def write_legend(wb: Workbook):
    ws = wb.create_sheet("Légende lots & statuts", 4)
    headers = ["Lot", "Nom", "Couverture principale (parcours)", "DoD GO"]
    for i, h in enumerate(headers, 1):
        ws.cell(1, i, h)
    style_header(ws, 4)

    rows = [
        ("0", "Fondations", P01, "Health + guards auth"),
        ("1", "Accès", f"{P02} · {P03}", "Login, invitation, reset, page publique"),
        ("2", "Clients magasin", P04, "CRUD + doublon + isolation"),
        ("3", "Paires & codes", P05, "À graver → Active, code immuable"),
        ("4", "Restitution magasin", P06, "Trouvé → remise sans PII si non client"),
        ("5", "Espace client", P07, "Liste ordonnée, perte, inactive"),
        ("6", "Notifications", P08, "Journal + inbox scoped"),
        ("7", "Envoi à distance", P09, "Stripe OK/KO + bordereau PII minimisée"),
        ("8", "Admin plateforme", P10, "KPIs + journal + blocage"),
        ("9", "Gravure & lots", P11, "CSV uniques, jamais recyclés"),
        ("10", "Facturation", P12, "Tarifs + Qonto skipped/queued"),
        ("—", "Sécurité", P13, "IDOR / XSS / audit / secrets"),
        ("—", "E2E référence", P14, "Parcours magasin complet"),
    ]
    for i, row in enumerate(rows, 2):
        for c, v in enumerate(row, 1):
            cell = ws.cell(i, c, v)
            cell.font = BODY_FONT
            cell.border = THIN
            cell.alignment = WRAP
        if i % 2 == 0:
            for c in range(1, 5):
                ws.cell(i, c).fill = ALT_FILL

    r = len(rows) + 3
    ws.cell(r, 1, "Statuts paire (Figma)").font = SECTION_FONT
    statuts = [
        ("a_graver", "À graver — code généré, pas encore marqué gravé"),
        ("active", "Active"),
        ("perdue", "Perdue (déclarée client)"),
        ("inactive", "Inactive (irréversible côté client)"),
        ("a_recuperer", "À récupérer en magasin dépositaire"),
        ("livraison_demandee", "Livraison demandée (après paiement OK)"),
        ("en_livraison", "En livraison (après prise en compte opticien)"),
        ("archivee", "Archivée / supprimée — code non réutilisable"),
    ]
    ws.cell(r + 1, 1, "Code statut").font = HEADER_FONT
    ws.cell(r + 1, 1).fill = HEADER_FILL
    ws.cell(r + 1, 2, "Libellé / règle").font = HEADER_FONT
    ws.cell(r + 1, 2).fill = HEADER_FILL
    for i, (code, label) in enumerate(statuts, r + 2):
        ws.cell(i, 1, code).font = BODY_FONT
        ws.cell(i, 2, label).font = BODY_FONT
        ws.cell(i, 1).border = THIN
        ws.cell(i, 2).border = THIN

    r2 = r + 2 + len(statuts) + 1
    ws.cell(r2, 1, "Deux codes").font = SECTION_FONT
    ws.cell(r2 + 1, 1, "Code assigné").font = Font(name="Calibri", bold=True)
    ws.cell(r2 + 1, 2, "FR + 4 chiffres + 2 lettres (ex. FR5487OK). Gravé / saisi. Jamais réutilisé.").font = BODY_FONT
    ws.cell(r2 + 2, 1, "Code restitution").font = Font(name="Calibri", bold=True)
    ws.cell(r2 + 2, 2, "Code temporaire pour remise magasin (email/SMS). ≠ code assigné. Rate-limit renvoi.").font = BODY_FONT

    autosize(ws, {"A": 18, "B": 48, "C": 48, "D": 40})


def main():
    wb = Workbook()
    # remove default
    default = wb.active
    wb.remove(default)

    write_mode_emploi(wb)
    write_parcours(wb)
    write_tests(wb)
    write_synthese(wb)
    write_legend(wb)

    wb.save(OUT)
    print(f"OK → {OUT}")
    print(f"Étapes tests : {len(TESTS)}")
    print(f"Parcours : {len(PARCOURS_META)}")


if __name__ == "__main__":
    main()
